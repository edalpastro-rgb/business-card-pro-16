# axon_knowledge_base.py — V50.z536 — Knowledge Base PERSONALE per utente (RAG locale, zero costi API per gli embedding)
#   V50.z536: multi-formato — PDF (anche scansionati → OCR), DOCX, XLSX, TXT/MD/CSV, immagini (OCR Tesseract locale ita+eng).
#
# • Vettori: ChromaDB embedded (PersistentClient su AXON_KB_DIR, volume `core_kb` sul VPS). MULTI-TENANT RIGIDO: una collection
#   per utente (`kb_<user_id>`), quindi i frammenti di un utente non possono MAI comparire nelle ricerche di un altro.
# • Embedding: all-MiniLM-L6-v2 in formato ONNX (onnxruntime + tokenizers, già nel Core; niente PyTorch, niente API esterne).
#   Modello in AXON_EMBED_MODEL_DIR (Dockerfile → /opt/axon-models/minilm; dev → backend/models/minilm). 384 dimensioni, cosine.
# • Chunking: finestre da KB_CHUNK_CHARS (800) con sovrapposizione KB_CHUNK_OVERLAP (100 ≈ 12,5 %), taglio su fine frase/paragrafo.
# • Ricerca: `_build_personal_kb_section(user_id, msg)` → i 3 frammenti più vicini (cosine ≥ soglia) come sezione aggiuntiva del
#   system prompt. Idle-summary e sliding window (21_axon_streaming) restano INTATTI: la sezione si aggiunge a `kb_section`.
# • Limiti per utente: 20 MB per file, 50 documenti, 5.000 frammenti totali. Indicizzazione in background (status: indexing → ready).
import asyncio
import base64
import os
import re
import threading
import uuid
from datetime import datetime
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
from fastapi import Depends, HTTPException
from pydantic import BaseModel, Field

KB_CHUNK_CHARS = 800
KB_CHUNK_OVERLAP = 100
KB_MAX_FILE_BYTES = 20 * 1024 * 1024
KB_MAX_DOCS = 50
KB_MAX_CHUNKS = 5000
KB_MAX_PDF_PAGES = 800
KB_TOP_K = 3
KB_MIN_SIMILARITY = 0.38
KB_EMBED_BATCH = 32
KB_EMBED_MAX_TOKENS = 256

def _kb_base_dir() -> str:
    # __file__ è server.py quando i moduli vengono exec()-uiti dal loader (→ backend/); routes/… se importato direttamente
    base = os.path.dirname(os.path.abspath(__file__))
    return os.path.dirname(base) if os.path.basename(base) == "routes" else base


_KB_DIR = os.environ.get("AXON_KB_DIR") or os.path.join(_kb_base_dir(), "data", "chroma")
_KB_MODEL_DIR = os.environ.get("AXON_EMBED_MODEL_DIR") or os.path.join(_kb_base_dir(), "models", "minilm")
_kb_log = logging.getLogger("axon.kb")  # noqa: F821 — logging importato in 00_core


# ───────────────────────────── Embedding locale (ONNX) ─────────────────────────────
class _LocalEmbedder:
    def __init__(self, model_dir: str):
        self.model_dir = model_dir
        self._lock = threading.Lock()
        self._session = None
        self._tok = None

    @property
    def available(self) -> bool:
        return os.path.isfile(os.path.join(self.model_dir, "model.onnx")) and os.path.isfile(os.path.join(self.model_dir, "tokenizer.json"))

    def _load(self):
        if self._session is not None:
            return
        with self._lock:
            if self._session is not None:
                return
            import onnxruntime as ort
            from tokenizers import Tokenizer
            tok = Tokenizer.from_file(os.path.join(self.model_dir, "tokenizer.json"))
            tok.enable_truncation(KB_EMBED_MAX_TOKENS)
            tok.enable_padding()
            opts = ort.SessionOptions()
            opts.intra_op_num_threads = max(1, min(4, os.cpu_count() or 1))
            self._session = ort.InferenceSession(os.path.join(self.model_dir, "model.onnx"), sess_options=opts, providers=["CPUExecutionProvider"])
            self._tok = tok
            _kb_log.info("[axon-kb] embedder MiniLM-L6-v2 ONNX caricato da %s", self.model_dir)

    def embed(self, texts: List[str]) -> List[List[float]]:
        if not texts:
            return []
        self._load()
        out: List[List[float]] = []
        for i in range(0, len(texts), KB_EMBED_BATCH):
            enc = self._tok.encode_batch(texts[i:i + KB_EMBED_BATCH])
            ids = np.array([e.ids for e in enc], dtype=np.int64)
            am = np.array([e.attention_mask for e in enc], dtype=np.int64)
            tt = np.array([e.type_ids for e in enc], dtype=np.int64)
            hidden = self._session.run(None, {"input_ids": ids, "attention_mask": am, "token_type_ids": tt})[0]
            mask = am[..., None].astype(np.float32)
            emb = (hidden * mask).sum(1) / np.clip(mask.sum(1), 1e-9, None)
            emb = emb / np.clip(np.linalg.norm(emb, axis=1, keepdims=True), 1e-9, None)
            out.extend(emb.astype(np.float32).tolist())
        return out


_kb_embedder = _LocalEmbedder(_KB_MODEL_DIR)
_kb_chroma_client = None
_kb_chroma_lock = threading.Lock()


def _kb_chroma():
    global _kb_chroma_client
    if _kb_chroma_client is None:
        with _kb_chroma_lock:
            if _kb_chroma_client is None:
                import chromadb
                from chromadb.config import Settings
                os.makedirs(_KB_DIR, exist_ok=True)
                _kb_chroma_client = chromadb.PersistentClient(path=_KB_DIR, settings=Settings(anonymized_telemetry=False, allow_reset=False))
                _kb_log.info("[axon-kb] ChromaDB persistente in %s", _KB_DIR)
    return _kb_chroma_client


def _kb_collection_name(user_id: str) -> str:
    safe = re.sub(r"[^a-zA-Z0-9_-]", "_", user_id or "anon")[:50]
    return f"kb_{safe}"


def _kb_collection(user_id: str):
    return _kb_chroma().get_or_create_collection(name=_kb_collection_name(user_id), metadata={"hnsw:space": "cosine"}, embedding_function=None)


# ───────────────────────────── Chunking ─────────────────────────────
_KB_BREAKS = ("\n\n", "\n", ". ", "! ", "? ", "; ", ", ", " ")


def _kb_chunk_text(text: str, size: int = KB_CHUNK_CHARS, overlap: int = KB_CHUNK_OVERLAP) -> List[str]:
    text = re.sub(r"[ \t]+", " ", text.replace("\r\n", "\n").replace("\r", "\n"))
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []
    chunks: List[str] = []
    start, n = 0, len(text)
    while start < n:
        end = min(n, start + size)
        if end < n:
            window = text[start:end]
            cut = -1
            for br in _KB_BREAKS:  # taglia sul separatore più "forte" presente nell'ultimo 30 % della finestra
                pos = window.rfind(br, int(size * 0.7))
                if pos > 0:
                    cut = pos + len(br)
                    break
            if cut > 0:
                end = start + cut
        piece = text[start:end].strip()
        if piece:
            chunks.append(piece)
        if end >= n:
            break
        start = max(end - overlap, start + 1)
    return chunks


# ───────────────────────────── V50.z536 — Multi-formato: DOCX · XLSX · immagini (OCR Tesseract locale) · PDF scansionati ─────────────────────────────
# OCR = binario `tesseract` di sistema (Dockerfile: tesseract-ocr + tesseract-ocr-ita), invocato via subprocess su PNG in memoria:
# nessuna dipendenza pip aggiuntiva, lingue ita+eng. I PDF senza layer testo (scansioni) vengono rasterizzati con PyMuPDF a 200 dpi.
KB_OCR_LANGS = os.environ.get("AXON_OCR_LANGS", "ita+eng")
KB_OCR_MAX_PAGES = int(os.environ.get("AXON_OCR_MAX_PAGES", "80"))   # pagine OCR per PDF scansionato (≈ 2-4 s/pagina su CPU)
KB_OCR_DPI = 200
KB_OCR_PAGE_TIMEOUT = 120
KB_PDF_TEXT_MIN_CHARS_PER_PAGE = 40  # sotto questa media il PDF è considerato "scansione" → OCR
KB_CAMERA_MAX_PAGES = 30             # V50.z537 — foto per singolo documento scattate dalla fotocamera
_KB_IMAGE_EXT = (".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp", ".gif", ".heic", ".heif")


def _kb_ocr_available() -> bool:
    import shutil
    return shutil.which("tesseract") is not None


def _kb_sniff_kind(filename: str, mime: str, raw: bytes) -> str:
    """→ 'pdf' | 'docx' | 'xlsx' | 'image' | 'txt' (magic bytes prima, poi estensione/mime)."""
    name = (filename or "").lower()
    m = (mime or "").lower()
    if raw[:5] == b"%PDF-" or name.endswith(".pdf") or m == "application/pdf":
        return "pdf"
    if raw[:4] == b"PK\x03\x04" or name.endswith((".docx", ".xlsx", ".xlsm")) or "officedocument" in m:
        if name.endswith((".xlsx", ".xlsm")) or "spreadsheet" in m:
            return "xlsx"
        if name.endswith(".docx") or "wordprocessingml" in m:
            return "docx"
        # ZIP generico: ispeziona il contenuto
        try:
            import zipfile
            names = zipfile.ZipFile(BytesIO(raw)).namelist()
            if any(n.startswith("xl/") for n in names):
                return "xlsx"
            if any(n.startswith("word/") for n in names):
                return "docx"
        except Exception:  # noqa: BLE001
            pass
    if m.startswith("image/") or name.endswith(_KB_IMAGE_EXT) or raw[:8] == b"\x89PNG\r\n\x1a\n" or raw[:3] == b"\xff\xd8\xff" or raw[:4] == b"RIFF":
        return "image"
    return "txt"


def _kb_ocr_pil(img) -> str:
    """OCR di un'immagine Pillow con tesseract (stdin PNG → stdout testo)."""
    import subprocess
    from PIL import ImageOps
    g = ImageOps.grayscale(img)
    if g.width < 1000:  # upscaling leggero: tesseract rende meglio con ≥ 300 dpi equivalenti
        f = min(3.0, 1000 / max(1, g.width))
        g = g.resize((int(g.width * f), int(g.height * f)))
    buf = BytesIO()
    g.save(buf, format="PNG")
    proc = subprocess.run(["tesseract", "stdin", "stdout", "-l", KB_OCR_LANGS, "--psm", "3"], input=buf.getvalue(), capture_output=True, timeout=KB_OCR_PAGE_TIMEOUT)
    if proc.returncode != 0:
        raise RuntimeError((proc.stderr or b"").decode("utf-8", errors="ignore")[:200] or "tesseract error")
    return proc.stdout.decode("utf-8", errors="ignore").strip()


def _kb_ocr_image_bytes(raw: bytes) -> str:
    from PIL import Image
    img = Image.open(BytesIO(raw))
    img.load()
    return _kb_ocr_pil(img.convert("RGB"))


def _kb_ocr_pdf(raw: bytes) -> Tuple[str, int]:
    """PDF scansionato → rasterizza ogni pagina (PyMuPDF) e la passa a tesseract. → (testo, pagine processate)."""
    import pymupdf as fitz  # PyMuPDF
    from PIL import Image
    doc = fitz.open(stream=raw, filetype="pdf")
    parts: List[str] = []
    n = min(len(doc), KB_OCR_MAX_PAGES)
    zoom = KB_OCR_DPI / 72.0
    for i in range(n):
        pix = doc[i].get_pixmap(matrix=fitz.Matrix(zoom, zoom), colorspace=fitz.csGRAY, alpha=False)
        img = Image.frombytes("L", (pix.width, pix.height), pix.samples)
        try:
            t = _kb_ocr_pil(img)
        except Exception as e:  # noqa: BLE001
            _kb_log.warning("[axon-kb] OCR pagina %d fallito: %s", i + 1, e)
            t = ""
        if t:
            parts.append(t)
    return "\n\n".join(parts), n


def _kb_pdf_text(raw: bytes) -> Tuple[str, int]:
    from pypdf import PdfReader
    reader = PdfReader(BytesIO(raw))
    parts: List[str] = []
    pages = 0
    for page in reader.pages[:KB_MAX_PDF_PAGES]:
        pages += 1
        try:
            t = (page.extract_text() or "").strip()
        except Exception:  # noqa: BLE001
            t = ""
        if t:
            parts.append(t)
    return "\n\n".join(parts), pages


def _kb_docx_text(raw: bytes) -> str:
    import docx  # python-docx
    d = docx.Document(BytesIO(raw))
    parts: List[str] = [p.text.strip() for p in d.paragraphs if p.text and p.text.strip()]
    for tbl in d.tables:
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _kb_xlsx_text(raw: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    parts: List[str] = []
    for ws in wb.worksheets:
        parts.append(f"\n\n## Foglio: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            if any(cells):
                parts.append(" | ".join(cells).rstrip(" |"))
    return "\n".join(parts).strip()


def _kb_extract_text(filename: str, mime: str, raw: bytes) -> Tuple[str, str, int, bool]:
    """→ (testo, source_type, pagine, needs_ocr). Estrazione RAPIDA (sincrona): PDF con layer testo via pypdf, DOCX, XLSX, txt/md/csv.
    needs_ocr=True ⇒ immagine o PDF scansionato: l'OCR (lento) viene eseguito in background da `_kb_ocr_and_index`."""
    kind = _kb_sniff_kind(filename, mime, raw)
    if kind == "pdf":
        text, pages = _kb_pdf_text(raw)
        scanned = pages > 0 and len(text.strip()) < KB_PDF_TEXT_MIN_CHARS_PER_PAGE * pages
        return text, "pdf", pages, scanned
    if kind == "docx":
        return _kb_docx_text(raw), "docx", 0, False
    if kind == "xlsx":
        return _kb_xlsx_text(raw), "xlsx", 0, False
    if kind == "image":
        return "", "image", 1, True
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return raw.decode(enc), "txt", 0, False
        except Exception:  # noqa: BLE001
            continue
    return raw.decode("utf-8", errors="ignore"), "txt", 0, False


async def _kb_ocr_and_index(doc_id: str, user_id: str, source_type: str, raw: Any) -> None:
    """Background: OCR (immagine singola, PIÙ foto/pagine [list di bytes] o PDF scansionato) → aggiorna il documento → indicizza."""
    try:
        if isinstance(raw, list):  # V50.z537 — foto scattate dalla fotocamera: una pagina per immagine
            texts: List[str] = []
            for i, page_raw in enumerate(raw, 1):
                try:
                    t = await asyncio.to_thread(_kb_ocr_image_bytes, page_raw)
                except Exception as e:  # noqa: BLE001
                    _kb_log.warning("[axon-kb] OCR foto %d fallito: %s", i, e)
                    t = ""
                if t.strip():
                    texts.append(f"--- Pagina {i} ---\n{t.strip()}" if len(raw) > 1 else t.strip())
                await db.kb_documents.update_one({"doc_id": doc_id}, {"$set": {"ocr_progress": f"{i}/{len(raw)}"}})  # noqa: F821
            text, pages = "\n\n".join(texts), len(raw)
        elif source_type == "image":
            text, pages = await asyncio.to_thread(_kb_ocr_image_bytes, raw), 1
        else:
            text, pages = await asyncio.to_thread(_kb_ocr_pdf, raw)
        if not text.strip():
            await db.kb_documents.update_one({"doc_id": doc_id}, {"$set": {"status": "error", "stage": None, "error": "L'OCR non ha riconosciuto testo leggibile (immagine troppo piccola o sfocata?)"}})  # noqa: F821
            return
        await db.kb_documents.update_one({"doc_id": doc_id}, {"$set": {"stage": "embedding", "ocr": True, "pages": pages, "chars": len(text)}})  # noqa: F821
        await _kb_index_document(doc_id, user_id, text)
    except Exception as e:  # noqa: BLE001
        _kb_log.exception("[axon-kb] OCR fallito per %s", doc_id)
        await db.kb_documents.update_one({"doc_id": doc_id}, {"$set": {"status": "error", "stage": None, "error": f"OCR non riuscito: {str(e)[:200]}"}})  # noqa: F821


# ───────────────────────────── Indicizzazione ─────────────────────────────
async def _kb_usage(user_id: str) -> Dict[str, int]:
    docs = await db.kb_documents.count_documents({"user_id": user_id})  # noqa: F821
    agg = await db.kb_documents.aggregate([{"$match": {"user_id": user_id}}, {"$group": {"_id": None, "chunks": {"$sum": "$chunks"}}}]).to_list(1)  # noqa: F821
    return {"docs": docs, "chunks": int(agg[0]["chunks"]) if agg else 0}


async def _kb_index_document(doc_id: str, user_id: str, text: str) -> None:
    try:
        chunks = _kb_chunk_text(text)
        usage = await _kb_usage(user_id)
        room = KB_MAX_CHUNKS - usage["chunks"]
        truncated = False
        if len(chunks) > room:
            chunks, truncated = chunks[:max(0, room)], True
        if not chunks:
            await db.kb_documents.update_one({"doc_id": doc_id}, {"$set": {"status": "error", "error": "Nessun testo estraibile (o limite frammenti raggiunto)", "chunks": 0}})  # noqa: F821
            return
        embeddings = await asyncio.to_thread(_kb_embedder.embed, chunks)
        col = _kb_collection(user_id)
        ids = [f"{doc_id}:{i}" for i in range(len(chunks))]
        metas = [{"doc_id": doc_id, "user_id": user_id, "idx": i} for i in range(len(chunks))]
        for i in range(0, len(ids), 500):  # batch: Chroma ha un limite per singola add
            await asyncio.to_thread(col.add, ids=ids[i:i + 500], embeddings=embeddings[i:i + 500], documents=chunks[i:i + 500], metadatas=metas[i:i + 500])
        await db.kb_documents.update_one({"doc_id": doc_id}, {"$set": {"status": "ready", "stage": None, "chunks": len(chunks), "chars": sum(len(c) for c in chunks), "truncated": truncated, "indexed_at": utcnow()}})  # noqa: F821
        _kb_log.info("[axon-kb] doc %s indicizzato: %d frammenti (utente %s)", doc_id, len(chunks), user_id)
    except Exception as e:  # noqa: BLE001
        _kb_log.exception("[axon-kb] indicizzazione fallita per %s", doc_id)
        await db.kb_documents.update_one({"doc_id": doc_id}, {"$set": {"status": "error", "error": str(e)[:300]}})  # noqa: F821


# ───────────────────────────── Ricerca (usata dal prompt Axon) ─────────────────────────────
async def _kb_personal_search(user_id: str, query: str, k: int = KB_TOP_K) -> List[Dict[str, Any]]:
    if not query or len(query.strip()) < 3 or not _kb_embedder.available:
        return []
    if await db.kb_documents.count_documents({"user_id": user_id, "status": "ready"}) == 0:  # noqa: F821
        return []
    qv = (await asyncio.to_thread(_kb_embedder.embed, [query.strip()[:2000]]))[0]
    col = _kb_collection(user_id)
    res = await asyncio.to_thread(col.query, query_embeddings=[qv], n_results=k, include=["documents", "metadatas", "distances"])
    docs = (res.get("documents") or [[]])[0]
    metas = (res.get("metadatas") or [[]])[0]
    dists = (res.get("distances") or [[]])[0]
    out: List[Dict[str, Any]] = []
    for d, m, dist in zip(docs, metas, dists):
        sim = 1.0 - float(dist)
        if sim >= KB_MIN_SIMILARITY:
            out.append({"text": d, "doc_id": (m or {}).get("doc_id"), "idx": (m or {}).get("idx"), "similarity": round(sim, 3)})
    if out:
        titles = {r["doc_id"]: r["title"] async for r in db.kb_documents.find({"doc_id": {"$in": list({o["doc_id"] for o in out})}}, {"_id": 0, "doc_id": 1, "title": 1})}  # noqa: F821
        for o in out:
            o["title"] = titles.get(o["doc_id"], "Documento")
    return out


async def _build_personal_kb_section(user_id: str, msg: str) -> str:
    hits = await _kb_personal_search(user_id, msg)
    if not hits:
        return ""
    lines = ["\n\nKNOWLEDGE BASE PERSONALE DELL'UTENTE (libri, appunti e documenti caricati da lui — frammenti più pertinenti alla domanda; "
             "usali come fonte quando rispondi e cita il titolo del documento se utile):"]
    for i, h in enumerate(hits, 1):
        lines.append(f"[{i}] «{h['title']}» (frammento {int(h['idx'] or 0) + 1}):\n{h['text'].strip()}")
    return "\n\n".join(lines) + "\n"


# ───────────────────────────── API ─────────────────────────────
class KbDocumentIn(BaseModel):
    title: Optional[str] = Field(default=None, max_length=200)
    text: Optional[str] = Field(default=None, max_length=2_000_000)  # testo incollato
    filename: Optional[str] = Field(default=None, max_length=255)
    mime: Optional[str] = None
    data_base64: Optional[str] = None  # PDF / DOCX / XLSX / TXT / immagine
    pages_base64: Optional[List[str]] = Field(default=None, max_length=KB_CAMERA_MAX_PAGES)  # V50.z537 — foto dalla fotocamera (una per pagina) → OCR


def _kb_b64_decode(s: str) -> bytes:
    s = s.split(",", 1)[1] if s.startswith("data:") else s
    return base64.b64decode(s)


def _kb_doc_out(d: dict) -> dict:
    for k in ("created_at", "indexed_at"):
        if isinstance(d.get(k), datetime):
            d[k] = d[k].isoformat()
    return d


@api_router.get("/kb/status")  # noqa: F821
async def kb_status(user=Depends(get_current_user)):  # noqa: F821
    usage = await _kb_usage(user["user_id"])
    return {"embedder_ready": _kb_embedder.available, "model": "all-MiniLM-L6-v2 (ONNX, locale)", "vector_db": "ChromaDB", "kb_dir": _KB_DIR,
            "ocr_ready": _kb_ocr_available(), "ocr_engine": "Tesseract (locale)", "ocr_langs": KB_OCR_LANGS, "formats": ["pdf", "docx", "xlsx", "txt", "md", "csv", "png", "jpg", "webp", "tiff"],
            "usage": usage, "limits": {"max_file_mb": KB_MAX_FILE_BYTES // (1024 * 1024), "max_docs": KB_MAX_DOCS, "max_chunks": KB_MAX_CHUNKS, "ocr_max_pages": KB_OCR_MAX_PAGES},
            "chunking": {"chars": KB_CHUNK_CHARS, "overlap": KB_CHUNK_OVERLAP}, "top_k": KB_TOP_K}


@api_router.get("/kb/documents")  # noqa: F821
async def kb_list(user=Depends(get_current_user)):  # noqa: F821
    rows = await db.kb_documents.find({"user_id": user["user_id"]}, {"_id": 0}).sort("created_at", -1).to_list(KB_MAX_DOCS + 5)  # noqa: F821
    usage = await _kb_usage(user["user_id"])
    return {"items": [_kb_doc_out(r) for r in rows], "usage": usage, "limits": {"max_file_mb": KB_MAX_FILE_BYTES // (1024 * 1024), "max_docs": KB_MAX_DOCS, "max_chunks": KB_MAX_CHUNKS}, "embedder_ready": _kb_embedder.available, "ocr_ready": _kb_ocr_available()}


@api_router.post("/kb/documents")  # noqa: F821
async def kb_add(body: KbDocumentIn, user=Depends(get_current_user)):  # noqa: F821
    uid = user["user_id"]
    if not _kb_embedder.available:
        raise HTTPException(status_code=503, detail="Motore di embedding locale non disponibile sul server (modello MiniLM mancante)")
    usage = await _kb_usage(uid)
    if usage["docs"] >= KB_MAX_DOCS:
        raise HTTPException(status_code=409, detail=f"Hai raggiunto il limite di {KB_MAX_DOCS} documenti: elimina qualcosa prima di caricare altro")
    if usage["chunks"] >= KB_MAX_CHUNKS:
        raise HTTPException(status_code=409, detail=f"Hai raggiunto il limite di {KB_MAX_CHUNKS} frammenti: elimina qualche documento")
    if body.pages_base64:
        # V50.z537 — foto scattate dalla fotocamera (1..N pagine): OCR in background, un unico documento
        if not _kb_ocr_available():
            raise HTTPException(status_code=400, detail="Le foto richiedono l'OCR ma Tesseract non è installato sul server")
        try:
            raw = [r for r in (_kb_b64_decode(p) for p in body.pages_base64 if p) if len(r) > 100]
        except Exception as e:  # noqa: BLE001
            raise HTTPException(status_code=400, detail=f"Foto non valida: {e}")
        if not raw:
            raise HTTPException(status_code=400, detail="Nessuna foto ricevuta")
        size = sum(len(r) for r in raw)
        if size > KB_MAX_FILE_BYTES:
            raise HTTPException(status_code=413, detail=f"Foto troppo pesanti nel complesso (max {KB_MAX_FILE_BYTES // (1024 * 1024)} MB)")
        text, source_type, pages, needs_ocr = "", "image", len(raw), True
    elif body.data_base64:
        try:
            raw = _kb_b64_decode(body.data_base64)
        except Exception as e:  # noqa: BLE001
            raise HTTPException(status_code=400, detail=f"File non valido: {e}")
        if len(raw) > KB_MAX_FILE_BYTES:
            raise HTTPException(status_code=413, detail=f"File troppo grande (max {KB_MAX_FILE_BYTES // (1024 * 1024)} MB)")
        try:
            text, source_type, pages, needs_ocr = await asyncio.to_thread(_kb_extract_text, body.filename or "", body.mime or "", raw)
        except Exception as e:  # noqa: BLE001
            raise HTTPException(status_code=400, detail=f"Impossibile leggere il file: {str(e)[:120]}")
        size = len(raw)
        if needs_ocr and not _kb_ocr_available():
            if text.strip():
                needs_ocr = False  # PDF con poco testo ma OCR non disponibile: usa quello che c'è
            else:
                raise HTTPException(status_code=400, detail="Questo file richiede l'OCR (immagine o PDF scansionato) ma Tesseract non è installato sul server")
    elif body.text and body.text.strip():
        text, source_type, pages, size, needs_ocr = body.text, "text", 0, len(body.text.encode("utf-8")), False
        raw = b""
    else:
        raise HTTPException(status_code=400, detail="Carica un file (PDF, DOCX, XLSX, TXT, immagine) oppure incolla del testo")
    if not needs_ocr and not text.strip():
        raise HTTPException(status_code=400, detail="Il documento non contiene testo leggibile")
    title = (body.title or "").strip() or (body.filename or "").rsplit(".", 1)[0] or text.strip().split("\n", 1)[0][:60] or ("Foto documento" if body.pages_base64 else "Documento")
    doc_id = f"kb_{uuid.uuid4().hex[:12]}"
    doc = {"doc_id": doc_id, "user_id": uid, "tenant_id": tenant_id(user), "title": title[:200], "source_type": source_type, "filename": body.filename,  # noqa: F821
           "size_bytes": size, "pages": pages, "chars": len(text), "chunks": 0, "status": "indexing", "stage": "ocr" if needs_ocr else "embedding", "ocr": bool(needs_ocr), "created_at": utcnow()}  # noqa: F821
    await db.kb_documents.insert_one(dict(doc))  # noqa: F821
    if needs_ocr:
        asyncio.create_task(_kb_ocr_and_index(doc_id, uid, source_type, raw))
    else:
        asyncio.create_task(_kb_index_document(doc_id, uid, text))
    return {"ok": True, "document": _kb_doc_out(doc), "ocr": bool(needs_ocr), "estimated_chunks": max(1, len(text) // (KB_CHUNK_CHARS - KB_CHUNK_OVERLAP)) if text else 0}


@api_router.delete("/kb/documents/{doc_id}")  # noqa: F821
async def kb_delete(doc_id: str, user=Depends(get_current_user)):  # noqa: F821
    uid = user["user_id"]
    d = await db.kb_documents.find_one({"doc_id": doc_id, "user_id": uid}, {"_id": 0, "doc_id": 1})  # noqa: F821
    if not d:
        raise HTTPException(status_code=404, detail="Documento non trovato")
    try:
        await asyncio.to_thread(_kb_collection(uid).delete, where={"doc_id": doc_id})
    except Exception as e:  # noqa: BLE001
        _kb_log.warning("[axon-kb] delete chroma %s: %s", doc_id, e)
    await db.kb_documents.delete_one({"doc_id": doc_id, "user_id": uid})  # noqa: F821
    return {"ok": True, "doc_id": doc_id}


class KbSearchIn(BaseModel):
    query: str = Field(min_length=2, max_length=2000)
    k: int = Field(default=KB_TOP_K, ge=1, le=10)


@api_router.post("/kb/search")  # noqa: F821
async def kb_search(body: KbSearchIn, user=Depends(get_current_user)):  # noqa: F821
    """Anteprima della ricerca (stessa funzione usata dal prompt di Axon)."""
    return {"hits": await _kb_personal_search(user["user_id"], body.query, body.k)}
