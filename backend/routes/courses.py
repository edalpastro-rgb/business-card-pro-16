# courses.py — V50.z275 — AREA CORSI (Corporate Gym v2.0 · Fase 1 core + Fase 2 test/relazione/dashboard + Fase 3 PowerPoint)
# Cartelle corsi create da Master/Manager (owner o utente con sotto-collaboratori), assegnate a singoli / tutti;
# contenuti (PDF, foto, audio, testo) salvati su Object Storage (fallback Mongo); tracciamento lettura (badge lampeggiante
# finché non aperta) e partecipazione (✅/❌ + motivo) con notifiche al creatore; lista partecipanti solo per il creatore;
# export Excel/PDF(HTML); reminder automatici (3/7/14 gg) via axon_notify (job giornaliero 08:45).
# Isolato: usa descendant_user_ids, tenant_id, axon_notify, _put_object_sync/_get_object_sync.

COURSE_REMINDER_DEFAULT_DAYS = 3
COURSE_PRE_REMINDER_DEFAULT = (1,)  # V50.z300 — Fase 5: promemoria 1 giorno prima della lezione
COURSE_PRE_REMINDER_ALLOWED = (0, 1, 2, 3, 7, 14)


def _course_pre_reminders(vals: Optional[List[int]]) -> List[int]:
    return sorted({int(v) for v in (vals or []) if int(v) in COURSE_PRE_REMINDER_ALLOWED}, reverse=True)


def _course_time(v: Optional[str]) -> str:
    m = re.match(r"^\s*(\d{1,2})[:.](\d{2})\s*$", v or "")
    if not m or int(m.group(1)) > 23 or int(m.group(2)) > 59:
        return ""
    return f"{int(m.group(1)):02d}:{m.group(2)}"


def _course_today_str() -> str:
    try:
        from zoneinfo import ZoneInfo
        return datetime.now(ZoneInfo("Europe/Rome")).strftime("%Y-%m-%d")
    except Exception:  # noqa: BLE001
        return utcnow().strftime("%Y-%m-%d")


class AttendanceIn(BaseModel):  # V50.z302 — spunta presenze del formatore
    present: List[str] = Field(default_factory=list)
    absent: List[str] = Field(default_factory=list)


class CourseFileIn(BaseModel):
    name: str
    mime: str = "application/octet-stream"
    data_base64: str


class CourseIn(BaseModel):
    title: str
    description: Optional[str] = ""
    body_text: Optional[str] = ""  # testo libero / appunti del corso
    course_date: Optional[str] = ""  # AAAA-MM-GG
    course_time: Optional[str] = ""  # V50.z300 — HH:MM (Fase 5)
    location: Optional[str] = ""
    maps_url: Optional[str] = ""
    recipients: List[str] = Field(default_factory=list)  # user_id; vuoto = tutti i collaboratori del creatore
    reminder_days: int = COURSE_REMINDER_DEFAULT_DAYS  # 0 = disattivato
    pre_reminders: List[int] = Field(default_factory=lambda: list(COURSE_PRE_REMINDER_DEFAULT))  # V50.z300 — giorni prima della lezione (0 = giorno stesso)
    files: List[CourseFileIn] = Field(default_factory=list)


class CourseUpdate(BaseModel):
    title: Optional[str] = None
    description: Optional[str] = None
    body_text: Optional[str] = None
    course_date: Optional[str] = None
    course_time: Optional[str] = None  # V50.z300
    location: Optional[str] = None
    maps_url: Optional[str] = None
    recipients: Optional[List[str]] = None
    reminder_days: Optional[int] = None
    pre_reminders: Optional[List[int]] = None  # V50.z300
    status: Optional[str] = None  # active | archived
    add_files: Optional[List[CourseFileIn]] = None
    remove_file_ids: Optional[List[str]] = None


class ParticipationIn(BaseModel):
    status: str  # yes | no
    notes: Optional[str] = ""


async def _course_can_manage(user: dict) -> bool:
    """Master (owner) o Manager = ha almeno un collaboratore sotto di sé."""
    if is_owner(user):
        return True
    return bool(await db.users.find_one({"parent_user_id": user["user_id"]}, {"_id": 1}))


async def _course_default_recipients(user: dict) -> List[str]:
    return await descendant_user_ids(user["user_id"])


async def _course_store_file(tid: str, course_id: str, f: CourseFileIn) -> dict:
    b64 = f.data_base64.split(",", 1)[1] if f.data_base64.startswith("data:") else f.data_base64
    if len(b64) > 24 * 1024 * 1024:
        raise HTTPException(status_code=413, detail=f"File troppo grande: {f.name}")
    fid = f"cf_{uuid.uuid4().hex[:10]}"
    meta = {"file_id": fid, "name": f.name[:160], "mime": f.mime or "application/octet-stream", "size": int(len(b64) * 0.75), "created_at": utcnow()}
    try:
        ext = (f.name.rsplit(".", 1)[-1] if "." in f.name else "bin")[:6]
        path = f"{_STORAGE_APP_NAME}/courses/{tid}/{course_id}/{fid}.{ext}"
        res = await run_in_threadpool(_put_object_sync, path, base64.b64decode(b64), meta["mime"])
        if res:
            meta["path"] = res.get("path") or path
            return meta
    except Exception as exc:  # noqa: BLE001
        logger.warning("[courses] object storage fallito (%s) → fallback Mongo", exc)
    meta["data_b64"] = b64
    return meta


def _course_public(c: dict, user: dict, names: Dict[str, str]) -> dict:
    uid = user["user_id"]
    is_creator = c.get("created_by") == uid
    reads = c.get("reads") or {}
    part = c.get("participation") or {}
    out = {k: v for k, v in c.items() if k not in ("_id", "reads", "participation")}
    out["files"] = [{k: v for k, v in f.items() if k != "data_b64"} for f in (c.get("files") or [])]
    out["is_creator"] = is_creator
    out["creator_name"] = names.get(c.get("created_by") or "", "")
    out["my_read_at"] = reads.get(uid)
    out["is_new"] = (not is_creator) and uid not in reads  # → badge lampeggiante
    mp = part.get(uid) or {}
    out["my_participation"] = mp.get("status") or "pending"
    out["my_notes"] = mp.get("notes") or ""
    out["recipient_count"] = len(c.get("recipients") or [])
    out["read_count"] = sum(1 for r in (c.get("recipients") or []) if r in reads)
    out["yes_count"] = sum(1 for r in (c.get("recipients") or []) if (part.get(r) or {}).get("status") == "yes")
    out["no_count"] = sum(1 for r in (c.get("recipients") or []) if (part.get(r) or {}).get("status") == "no")
    if is_creator:  # lista partecipanti: SOLO il creatore
        certs = c.get("certificates") or {}
        att = c.get("attendance") or {}
        out["participants"] = [{
            "user_id": r, "name": names.get(r, r), "read_at": reads.get(r), "status": (part.get(r) or {}).get("status") or "pending",
            "notes": (part.get(r) or {}).get("notes") or "", "answered_at": (part.get(r) or {}).get("timestamp"),
            "certificate_issued": bool(certs.get(r)), "certificate_id": (certs.get(r) or {}).get("cert_id"),  # V50.z300
            "present": (att.get(r) or {}).get("present") if r in att else None,  # V50.z302 — True/False/None (non registrata)
        } for r in (c.get("recipients") or [])]
        out["certificates_issued"] = sum(1 for r in (c.get("recipients") or []) if certs.get(r))
        out["present_count"] = sum(1 for r in (c.get("recipients") or []) if (att.get(r) or {}).get("present"))
    out["attendance_taken_at"] = c.get("attendance_taken_at")
    out["is_today_or_past"] = bool(c.get("course_date")) and c["course_date"] <= _course_today_str()
    return out


async def _course_names(ids: List[str]) -> Dict[str, str]:
    names: Dict[str, str] = {}
    async for u in db.users.find({"user_id": {"$in": list(set(ids))}}, {"_id": 0, "user_id": 1, "name": 1, "email": 1}):
        names[u["user_id"]] = u.get("name") or u.get("email") or u["user_id"]
    return names


async def _course_get(course_id: str, user: dict) -> dict:
    c = await db.courses_folders.find_one({"course_id": course_id, "tenant_id": tenant_id(user)}, {"_id": 0})
    if not c:
        raise HTTPException(status_code=404, detail="Corso non trovato")
    if c.get("created_by") != user["user_id"] and user["user_id"] not in (c.get("recipients") or []) and not is_owner(user):
        raise HTTPException(status_code=403, detail="Corso non assegnato a te")
    return c


@api_router.get("/courses/can-manage")
async def courses_can_manage(user=Depends(get_current_user)):
    team = await _course_default_recipients(user)
    names = await _course_names(team)
    return {"can_manage": await _course_can_manage(user), "team": [{"user_id": u, "name": names.get(u, u)} for u in team]}


@api_router.post("/courses")
async def course_create(body: CourseIn, user=Depends(get_current_user)):
    if not await _course_can_manage(user):
        raise HTTPException(status_code=403, detail="Solo Master o Manager possono creare cartelle corsi")
    tid = tenant_id(user)
    course_id = f"crs_{uuid.uuid4().hex[:12]}"
    recipients = body.recipients or await _course_default_recipients(user)
    files = [await _course_store_file(tid, course_id, f) for f in body.files[:20]]
    doc = {
        "course_id": course_id, "tenant_id": tid, "created_by": user["user_id"], "title": body.title.strip()[:160] or "Corso",
        "description": (body.description or "").strip()[:2000], "body_text": (body.body_text or "")[:20000], "course_date": (body.course_date or "")[:10],
        "course_time": _course_time(body.course_time), "pre_reminders": _course_pre_reminders(body.pre_reminders),
        "location": (body.location or "").strip()[:200], "maps_url": (body.maps_url or "").strip()[:500], "recipients": recipients,
        "reminder_days": max(0, min(30, int(body.reminder_days or 0))), "files": files, "status": "active", "reads": {}, "participation": {},
        "reminders_sent": {}, "created_at": utcnow(), "updated_at": utcnow(),
    }
    await db.courses_folders.insert_one(doc)
    for r in recipients:
        try:
            await axon_notify(r, tid, "reminder", f"📚 Nuovo corso: {doc['title']}", "Apri la cartella per leggere i materiali e indicare se partecipi.",
                              deeplink=f"/courses/{course_id}", dedupe_key=f"course_new_{course_id}_{r}", team_copy=False)
        except Exception:  # noqa: BLE001
            pass
    names = await _course_names([user["user_id"], *recipients])
    return _course_public(doc, user, names)


@api_router.get("/courses")
async def courses_list(include_archived: bool = False, user=Depends(get_current_user)):
    q: Dict[str, Any] = {"tenant_id": tenant_id(user), "$or": [{"created_by": user["user_id"]}, {"recipients": user["user_id"]}]}
    if is_owner(user):
        q.pop("$or")
    if not include_archived:
        q["status"] = "active"
    items = await db.courses_folders.find(q, {"_id": 0, "files.data_b64": 0, "body_text": 0}).sort("created_at", -1).to_list(300)
    ids = {c["created_by"] for c in items} | {r for c in items for r in (c.get("recipients") or [])}
    names = await _course_names(list(ids))
    out = [_course_public(c, user, names) for c in items]
    return {"items": out, "new_count": sum(1 for c in out if c["is_new"]), "can_manage": await _course_can_manage(user)}


@api_router.get("/courses/{course_id}")
async def course_get(course_id: str, user=Depends(get_current_user)):
    c = await _course_get(course_id, user)
    names = await _course_names([c["created_by"], *(c.get("recipients") or [])])
    out = _course_public(c, user, names)
    # V50.z276 — idoneità attestato (mio + per partecipante se creatore)
    t = await db.courses_tests.find_one({"course_id": course_id}, {"_id": 0, "questions": 1, "responses": 1})
    certs = c.get("certificates") or {}
    out["has_test"] = bool(t and t.get("questions"))
    out["my_certificate"] = {"eligible": _course_cert_eligible(c, t, user["user_id"]), **(certs.get(user["user_id"]) or {})}
    for p in out.get("participants") or []:
        p["certificate_eligible"] = _course_cert_eligible(c, t, p["user_id"])
    return out


@api_router.post("/courses/{course_id}/read")
async def course_mark_read(course_id: str, user=Depends(get_current_user)):
    """Primo click sulla cartella: ferma il lampeggio, registra timestamp, notifica il creatore."""
    c = await _course_get(course_id, user)
    uid = user["user_id"]
    if c.get("created_by") == uid or uid in (c.get("reads") or {}):
        return {"ok": True, "already": True}
    now = utcnow()
    await db.courses_folders.update_one({"course_id": course_id}, {"$set": {f"reads.{uid}": now}})
    try:
        await axon_notify(c["created_by"], c["tenant_id"], "reminder", f"👀 {user.get('name') or user.get('email')} ha letto il corso «{c['title']}»",
                          f"Letto il {now.astimezone().strftime('%d/%m/%Y %H:%M')}.", deeplink=f"/courses/{course_id}", dedupe_key=f"course_read_{course_id}_{uid}", team_copy=False)
    except Exception:  # noqa: BLE001
        pass
    return {"ok": True, "read_at": now}


@api_router.post("/courses/{course_id}/participation")
async def course_participation(course_id: str, body: ParticipationIn, user=Depends(get_current_user)):
    c = await _course_get(course_id, user)
    if body.status not in ("yes", "no"):
        raise HTTPException(status_code=400, detail="status deve essere yes|no")
    uid = user["user_id"]
    now = utcnow()
    rec = {"status": body.status, "notes": (body.notes or "").strip()[:1000], "timestamp": now}
    await db.courses_folders.update_one({"course_id": course_id}, {"$set": {f"participation.{uid}": rec, f"reads.{uid}": (c.get("reads") or {}).get(uid) or now}})
    who = user.get("name") or user.get("email")
    title = f"✅ {who}: Partecipo" if body.status == "yes" else f"❌ {who}: Non partecipo"
    msg = f"Corso «{c['title']}»" + (f" — motivo: {rec['notes']}" if body.status == "no" and rec["notes"] else "")
    try:
        await axon_notify(c["created_by"], c["tenant_id"], "reminder", title, msg, deeplink=f"/courses/{course_id}",
                          dedupe_key=f"course_part_{course_id}_{uid}_{body.status}_{now.strftime('%Y%m%d%H%M')}", extra={"urgency": "HIGH" if body.status == "no" else "NORMAL"}, team_copy=False)
    except Exception:  # noqa: BLE001
        pass
    if body.status == "yes":
        await _course_cert_issue_if_ready(c, uid)  # V50.z276
    return {"ok": True, "participation": rec}


@api_router.put("/courses/{course_id}")
async def course_update(course_id: str, body: CourseUpdate, user=Depends(get_current_user)):
    c = await _course_get(course_id, user)
    if c.get("created_by") != user["user_id"] and not is_owner(user):
        raise HTTPException(status_code=403, detail="Solo il creatore può modificare il corso")
    upd: Dict[str, Any] = {"updated_at": utcnow()}
    for k in ("title", "description", "body_text", "course_date", "location", "maps_url", "recipients", "status"):
        v = getattr(body, k)
        if v is not None:
            upd[k] = v.strip() if isinstance(v, str) and k != "body_text" else v
    if body.reminder_days is not None:
        upd["reminder_days"] = max(0, min(30, int(body.reminder_days)))
    if body.course_time is not None:
        upd["course_time"] = _course_time(body.course_time)
    if body.pre_reminders is not None:
        upd["pre_reminders"] = _course_pre_reminders(body.pre_reminders)
    files = list(c.get("files") or [])
    if body.remove_file_ids:
        files = [f for f in files if f.get("file_id") not in set(body.remove_file_ids)]
    for f in (body.add_files or [])[:20]:
        files.append(await _course_store_file(c["tenant_id"], course_id, f))
    upd["files"] = files
    await db.courses_folders.update_one({"course_id": course_id}, {"$set": upd})
    new_recipients = [r for r in (body.recipients or []) if r not in (c.get("recipients") or [])]
    for r in new_recipients:
        try:
            await axon_notify(r, c["tenant_id"], "reminder", f"📚 Nuovo corso: {upd.get('title') or c['title']}", "Apri la cartella per leggere i materiali e indicare se partecipi.",
                              deeplink=f"/courses/{course_id}", dedupe_key=f"course_new_{course_id}_{r}", team_copy=False)
        except Exception:  # noqa: BLE001
            pass
    c2 = await db.courses_folders.find_one({"course_id": course_id}, {"_id": 0})
    names = await _course_names([c2["created_by"], *(c2.get("recipients") or [])])
    return _course_public(c2, user, names)


@api_router.delete("/courses/{course_id}")
async def course_delete(course_id: str, user=Depends(get_current_user)):
    c = await _course_get(course_id, user)
    if c.get("created_by") != user["user_id"] and not is_owner(user):
        raise HTTPException(status_code=403, detail="Solo il creatore può eliminare il corso")
    await db.courses_folders.delete_one({"course_id": course_id})
    await db.courses_tests.delete_one({"course_id": course_id})
    return {"ok": True}


@api_router.get("/courses/{course_id}/files/{file_id}")
async def course_file(course_id: str, file_id: str, user=Depends(get_current_user)):
    c = await _course_get(course_id, user)
    f = next((x for x in (c.get("files") or []) if x.get("file_id") == file_id), None)
    if not f:
        raise HTTPException(status_code=404, detail="File non trovato")
    b64 = None
    if f.get("path"):
        got = await run_in_threadpool(_get_object_sync, f["path"])
        if got:
            b64 = base64.b64encode(got[0]).decode()
    b64 = b64 or f.get("data_b64")
    if not b64:
        raise HTTPException(status_code=404, detail="Contenuto non disponibile")
    return {"name": f["name"], "mime": f.get("mime") or "application/octet-stream", "data_base64": b64}


@api_router.get("/courses/{course_id}/export")
async def course_export(course_id: str, fmt: str = "xlsx", user=Depends(get_current_user)):
    """Riepilogo partecipanti (solo creatore/Master): xlsx (openpyxl) oppure html (stampabile in PDF)."""
    c = await _course_get(course_id, user)
    if c.get("created_by") != user["user_id"] and not is_owner(user):
        raise HTTPException(status_code=403, detail="Solo il creatore può esportare la lista")
    names = await _course_names(c.get("recipients") or [])
    reads, part = c.get("reads") or {}, c.get("participation") or {}
    rows = []
    for r in c.get("recipients") or []:
        p = part.get(r) or {}
        ra = reads.get(r)
        rows.append([names.get(r, r), ra.astimezone().strftime("%d/%m/%Y %H:%M") if isinstance(ra, datetime) else "non letto",
                     {"yes": "Partecipa", "no": "Non partecipa"}.get(p.get("status"), "In attesa"), p.get("notes") or ""])
    safe = re.sub(r"[^A-Za-z0-9_\- ]+", "", c["title"]).strip() or "corso"
    if fmt == "xlsx":
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Partecipanti"
        ws.append([f"Corso: {c['title']}", "", "", f"Data: {c.get('course_date') or '-'}"])
        ws.append(["Collaboratore", "Lettura", "Partecipazione", "Motivazioni"])
        for row in rows:
            ws.append(row)
        for col, w in zip("ABCD", (32, 20, 18, 60)):
            ws.column_dimensions[col].width = w
        buf = io.BytesIO(); wb.save(buf)
        return {"name": f"{safe}_partecipanti.xlsx", "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "data_base64": base64.b64encode(buf.getvalue()).decode()}
    trs = "".join(f"<tr><td>{html_lib.escape(a)}</td><td>{html_lib.escape(b)}</td><td>{html_lib.escape(cc)}</td><td>{html_lib.escape(d)}</td></tr>" for a, b, cc, d in rows)
    html = (f"<html><head><meta charset='utf-8'><style>body{{font-family:Helvetica,Arial;padding:24px}}table{{border-collapse:collapse;width:100%}}td,th{{border:1px solid #cbd5e1;padding:6px 8px;font-size:13px;text-align:left}}th{{background:#ede9fe}}</style></head>"
            f"<body><h2>Corso: {html_lib.escape(c['title'])}</h2><p>Data corso: {html_lib.escape(c.get('course_date') or '-')} · Letture {sum(1 for r in c.get('recipients') or [] if r in reads)}/{len(c.get('recipients') or [])}</p>"
            f"<table><tr><th>Collaboratore</th><th>Lettura</th><th>Partecipazione</th><th>Motivazioni</th></tr>{trs}</table></body></html>")
    return {"name": f"{safe}_partecipanti.html", "mime": "text/html", "html": html}


# ---------- Reminder automatici: non ha aperto entro X giorni / ha letto ma non ha risposto ----------
# V50.z300 — Fase 5: + promemoria PRIMA della lezione (pre_reminders: N giorni prima / giorno stesso) ai partecipanti
# (✅ o ⏳, mai a chi ha detto ❌) e al creatore (riepilogo presenze); il giorno DOPO la lezione: attestati emessi
# in automatico agli idonei, «completa il test» a chi manca, riepilogo attestati al creatore.
def _course_when_text(c: dict, days_until: int) -> str:
    when = "Oggi" if days_until == 0 else "Domani" if days_until == 1 else f"Tra {days_until} giorni"
    d = (c.get("course_date") or "").split("-")
    parts = [when + (" alle " + c["course_time"] if c.get("course_time") else "") + (" (" + "/".join(reversed(d)) + ")" if len(d) == 3 and days_until > 1 else "")]
    if c.get("location"):
        parts.append("📍 " + c["location"])
    return " · ".join(parts)


async def _courses_pre_lesson(c: dict, today, sent: dict, now: datetime) -> int:
    """Promemoria pre-lezione + post-lezione per un corso. Ritorna il numero di notifiche inviate."""
    n = 0
    try:
        from datetime import date as _date
        y, m, d = (c.get("course_date") or "").split("-")
        cdate = _date(int(y), int(m), int(d))
    except Exception:  # noqa: BLE001
        return 0
    days_until = (cdate - today).days
    part = c.get("participation") or {}
    recipients = c.get("recipients") or []
    tid, cid, title = c["tenant_id"], c["course_id"], c["title"]
    link = f"/courses/{cid}"

    async def _send(uid: str, key: str, ttl: str, msg: str, urgency: str = "NORMAL") -> None:
        nonlocal n
        if sent.get(key):
            return
        try:
            await axon_notify(uid, tid, "reminder", ttl, msg, deeplink=link, dedupe_key=f"course_rem_{cid}_{key}", extra={"urgency": urgency}, team_copy=False)
            sent[key] = now.isoformat(); n += 1
        except Exception:  # noqa: BLE001
            pass

    if days_until in (c.get("pre_reminders") or []):
        when = _course_when_text(c, days_until)
        for r in recipients:
            st = (part.get(r) or {}).get("status")
            if st == "no":
                continue
            msg = f"{when}. " + ("Ti aspettiamo!" if st == "yes" else "Non hai ancora indicato se partecipi: rispondi ✅ / ❌.")
            await _send(r, f"pre{days_until}_{r}", f"📅 Corso {'oggi' if days_until == 0 else 'domani' if days_until == 1 else f'tra {days_until} giorni'}: {title}", msg, "HIGH" if days_until <= 1 else "NORMAL")
        yes = sum(1 for r in recipients if (part.get(r) or {}).get("status") == "yes")
        no = sum(1 for r in recipients if (part.get(r) or {}).get("status") == "no")
        await _send(c["created_by"], f"pre{days_until}_creator", f"📅 Il tuo corso «{title}»: {when}",
                    f"Presenze: ✅ {yes} partecipano · ❌ {no} · ⏳ {len(recipients) - yes - no} in attesa di risposta.")
    if days_until == -1:  # giorno dopo la lezione
        t = await db.courses_tests.find_one({"course_id": cid}, {"_id": 0, "questions": 1, "responses": 1})
        issued, missing_test = 0, 0
        for r in recipients:
            if (part.get(r) or {}).get("status") != "yes":
                continue
            await _course_cert_issue_if_ready(c, r)  # notifica «🎓 Attestato disponibile» se idoneo e non ancora emesso
            fresh = await db.courses_folders.find_one({"course_id": cid}, {"_id": 0, "certificates": 1}) or {}
            if (fresh.get("certificates") or {}).get(r):
                issued += 1
            elif t and t.get("questions"):
                missing_test += 1
                await _send(r, f"post_test_{r}", f"📝 Completa il test: {title}", "Il corso si è svolto ieri: completa il test per ricevere il tuo attestato di partecipazione.")
        await _send(c["created_by"], "post_creator", f"🎓 Corso svolto: {title}",
                    f"Attestati pronti: {issued}" + (f" · {missing_test} partecipanti devono ancora completare il test" if missing_test else "") + ". Apri il corso per scaricarli o rilasciarli a tutti.")
    return n


async def _courses_reminder_job():
    n = 0
    now = utcnow()
    try:
        from zoneinfo import ZoneInfo
        today = datetime.now(ZoneInfo("Europe/Rome")).date()
    except Exception:  # noqa: BLE001
        today = now.date()
    try:
        tests_map = {t["course_id"]: t for t in await db.courses_tests.find({}, {"_id": 0, "course_id": 1, "responses": 1, "reminder_days": 1}).to_list(2000)}
        async for c in db.courses_folders.find({"status": "active"}, {"_id": 0, "files": 0, "body_text": 0}):
            sent = c.get("reminders_sent") or {}
            before = dict(sent)
            if c.get("course_date"):
                n += await _courses_pre_lesson(c, today, sent, now)
            days = int(c.get("reminder_days") or 0)
            created = c.get("created_at")
            has_test = c["course_id"] in tests_map
            if isinstance(created, datetime) and (days > 0 or has_test):
                created = created if created.tzinfo else created.replace(tzinfo=timezone.utc)
                if not (days > 0 and (now - created).days < days and not has_test):
                    reads, part = c.get("reads") or {}, c.get("participation") or {}
                    for r in c.get("recipients") or []:
                        key = None
                        if days > 0 and (now - created).days >= days and r not in reads:
                            key, title, msg = f"read_{r}", f"⏰ Corso da leggere: {c['title']}", f"Non hai ancora aperto la cartella del corso (da {days} giorni)."
                        elif not (part.get(r) or {}).get("status"):
                            key, title, msg = f"part_{r}", f"⏰ Indica se partecipi: {c['title']}", "Hai letto il corso ma non hai ancora risposto ✅ Partecipo / ❌ Non partecipo."
                        elif tests_map.get(c["course_id"]) and r not in (tests_map[c["course_id"]].get("responses") or {}) and int(tests_map[c["course_id"]].get("reminder_days") or 0) > 0 \
                                and (now - created).days >= int(tests_map[c["course_id"]].get("reminder_days") or 0):
                            key, title, msg = f"test_{r}", f"📝 Test da completare: {c['title']}", "Non hai ancora completato il test del corso."
                        if not key or sent.get(key):
                            continue
                        try:
                            await axon_notify(r, c["tenant_id"], "reminder", title, msg, deeplink=f"/courses/{c['course_id']}", dedupe_key=f"course_rem_{c['course_id']}_{key}", team_copy=False)
                            sent[key] = now.isoformat(); n += 1
                        except Exception:  # noqa: BLE001
                            pass
            if sent != before:
                await db.courses_folders.update_one({"course_id": c["course_id"]}, {"$set": {"reminders_sent": sent}})
    except Exception as exc:  # noqa: BLE001
        logger.warning("[courses] reminder job error: %s", exc)
    return {"sent": n}


@api_router.post("/courses/run-reminders")
async def courses_run_reminders(user=Depends(get_current_user)):
    if not await _course_can_manage(user):
        raise HTTPException(status_code=403, detail="Solo Master/Manager")
    return await _courses_reminder_job()


@api_router.post("/courses/{course_id}/certificates/issue-all")
async def course_certificates_issue_all(course_id: str, user=Depends(get_current_user)):
    """V50.z300 — Il creatore rilascia l'attestato a TUTTI i partecipanti ✅ (anche senza test completato) → notifica «🎓 Attestato pronto»."""
    c = await _course_creator_only(course_id, user)
    part, certs, att = c.get("participation") or {}, dict(c.get("certificates") or {}), c.get("attendance") or {}
    taken = bool(c.get("attendance_taken_at"))
    issued = []
    for r in c.get("recipients") or []:
        if (part.get(r) or {}).get("status") != "yes" or certs.get(r):
            continue
        if taken and not (att.get(r) or {}).get("present"):  # V50.z302 — presenze registrate: solo ai presenti
            continue
        cert = {"cert_id": f"ATT-{utcnow().strftime('%Y%m%d')}-{uuid.uuid4().hex[:6].upper()}", "issued_at": utcnow(), "issued_by": user["user_id"]}
        await db.courses_folders.update_one({"course_id": course_id}, {"$set": {f"certificates.{r}": cert}})
        certs[r] = cert; issued.append(r)
        try:
            await axon_notify(r, c["tenant_id"], "reminder", f"🎓 Attestato pronto: {c['title']}", "Il formatore ha rilasciato il tuo attestato di partecipazione: scaricalo in PDF dalla cartella del corso.",
                              deeplink=f"/courses/{course_id}", dedupe_key=f"course_cert_{course_id}_{r}", extra={"urgency": "HIGH"}, team_copy=False)
        except Exception:  # noqa: BLE001
            pass
    return {"ok": True, "issued": len(issued), "total_certificates": len(certs)}


@api_router.put("/courses/{course_id}/attendance")
async def course_attendance(course_id: str, body: AttendanceIn, user=Depends(get_current_user)):
    """V50.z302 — Spunta presenze del formatore (giorno del corso): presenti/assenti tra i destinatari. Da questo momento
    l'attestato spetta SOLO ai presenti (anche il rilascio automatico del giorno dopo e «Rilascia a tutti»)."""
    c = await _course_creator_only(course_id, user)
    recipients = set(c.get("recipients") or [])
    now = utcnow()
    upd: Dict[str, Any] = {"attendance_taken_at": now, "updated_at": now}
    for r in body.present:
        if r in recipients:
            upd[f"attendance.{r}"] = {"present": True, "at": now, "by": user["user_id"]}
    for r in body.absent:
        if r in recipients and r not in body.present:
            upd[f"attendance.{r}"] = {"present": False, "at": now, "by": user["user_id"]}
    await db.courses_folders.update_one({"course_id": course_id}, {"$set": upd})
    c2 = await db.courses_folders.find_one({"course_id": course_id}, {"_id": 0})
    names = await _course_names([c2["created_by"], *(c2.get("recipients") or [])])
    out = _course_public(c2, user, names)
    t = await db.courses_tests.find_one({"course_id": course_id}, {"_id": 0, "questions": 1, "responses": 1})
    out["has_test"] = bool(t and t.get("questions"))
    for p in out.get("participants") or []:
        p["certificate_eligible"] = _course_cert_eligible(c2, t, p["user_id"])
    return out


# =====================================================================================
# V50.z274 — Fase 2: TEST/QUIZ · RELAZIONE VOCALE FINE CORSO · DASHBOARD AZIENDA
# =====================================================================================
class QuizQuestionIn(BaseModel):
    id: Optional[str] = None
    text: str
    type: str = "mcq"  # mcq | open
    options: List[str] = Field(default_factory=list)
    correct: Optional[int] = None  # indice opzione corretta (mcq)


class QuizIn(BaseModel):
    questions: List[QuizQuestionIn]
    reminder_days: int = 3  # 0 = off


class QuizSubmitIn(BaseModel):
    answers: Dict[str, Any]  # question_id → indice (mcq) o testo (open)
    duration_s: Optional[int] = 0


class CourseReportIn(BaseModel):
    transcript: Optional[str] = None
    summary: Optional[str] = None
    highlights: Optional[List[str]] = None
    actions: Optional[List[str]] = None
    feedback: Optional[str] = None
    proposals: Optional[List[str]] = None
    visible_to_recipients: Optional[bool] = None


async def _course_creator_only(course_id: str, user: dict) -> dict:
    c = await _course_get(course_id, user)
    if c.get("created_by") != user["user_id"] and not is_owner(user):
        raise HTTPException(status_code=403, detail="Solo il creatore del corso")
    return c


def _quiz_public(t: dict, user: dict, creator: bool) -> dict:
    qs = []
    for q in t.get("questions") or []:
        qq = {"id": q["id"], "text": q["text"], "type": q["type"], "options": q.get("options") or []}
        if creator:
            qq["correct"] = q.get("correct")
        qs.append(qq)
    resp = (t.get("responses") or {})
    mine = resp.get(user["user_id"])
    out = {"course_id": t["course_id"], "questions": qs, "reminder_days": t.get("reminder_days", 3), "my_response": mine}
    if creator:
        n = len(resp)
        scored = [r for r in resp.values() if r.get("total")]
        out["stats"] = {"responses": n, "avg_score_pct": round(sum(r["score"] / r["total"] * 100 for r in scored) / len(scored), 1) if scored else None,
                        "avg_duration_s": round(sum(int(r.get("duration_s") or 0) for r in resp.values()) / n) if n else 0,
                        "by_user": {uid: {"score": r.get("score"), "total": r.get("total"), "timestamp": r.get("timestamp"), "open_answers": r.get("open_answers") or {}} for uid, r in resp.items()}}
    return out


@api_router.put("/courses/{course_id}/test")
async def course_test_upsert(course_id: str, body: QuizIn, user=Depends(get_current_user)):
    await _course_creator_only(course_id, user)
    qs = []
    for i, q in enumerate(body.questions[:40]):
        if not q.text.strip():
            continue
        typ = "open" if q.type == "open" else "mcq"
        opts = [o.strip() for o in q.options if o.strip()][:8]
        if typ == "mcq" and len(opts) < 2:
            raise HTTPException(status_code=400, detail=f"Domanda {i + 1}: servono almeno 2 opzioni")
        qs.append({"id": q.id or f"q_{uuid.uuid4().hex[:6]}", "text": q.text.strip()[:600], "type": typ, "options": opts, "correct": q.correct if typ == "mcq" and q.correct is not None and 0 <= int(q.correct) < len(opts) else None})
    await db.courses_tests.update_one({"course_id": course_id}, {"$set": {"course_id": course_id, "questions": qs, "reminder_days": max(0, min(30, int(body.reminder_days))), "updated_at": utcnow()}, "$setOnInsert": {"responses": {}, "created_at": utcnow()}}, upsert=True)
    t = await db.courses_tests.find_one({"course_id": course_id}, {"_id": 0})
    return _quiz_public(t, user, True)


@api_router.get("/courses/{course_id}/test")
async def course_test_get(course_id: str, user=Depends(get_current_user)):
    c = await _course_get(course_id, user)
    t = await db.courses_tests.find_one({"course_id": course_id}, {"_id": 0})
    if not t:
        return {"course_id": course_id, "questions": [], "my_response": None}
    creator = c.get("created_by") == user["user_id"] or is_owner(user)
    out = _quiz_public(t, user, creator)
    if creator and out.get("stats"):
        names = await _course_names(list(out["stats"]["by_user"].keys()))
        for uid, r in out["stats"]["by_user"].items():
            r["name"] = names.get(uid, uid)
    return out


@api_router.post("/courses/{course_id}/test/submit")
async def course_test_submit(course_id: str, body: QuizSubmitIn, user=Depends(get_current_user)):
    c = await _course_get(course_id, user)
    t = await db.courses_tests.find_one({"course_id": course_id}, {"_id": 0})
    if not t or not t.get("questions"):
        raise HTTPException(status_code=404, detail="Nessun test per questo corso")
    score = total = 0
    open_answers: Dict[str, str] = {}
    for q in t["questions"]:
        a = body.answers.get(q["id"])
        if q["type"] == "mcq":
            total += 1
            try:
                if a is not None and q.get("correct") is not None and int(a) == int(q["correct"]):
                    score += 1
            except Exception:  # noqa: BLE001
                pass
        elif a:
            open_answers[q["id"]] = str(a)[:2000]
    rec = {"answers": {k: v for k, v in body.answers.items()}, "open_answers": open_answers, "score": score, "total": total, "duration_s": int(body.duration_s or 0), "timestamp": utcnow()}
    await db.courses_tests.update_one({"course_id": course_id}, {"$set": {f"responses.{user['user_id']}": rec}})
    try:
        await axon_notify(c["created_by"], c["tenant_id"], "reminder", f"📝 {user.get('name') or user.get('email')} ha completato il test «{c['title']}»",
                          f"Punteggio {score}/{total}" + (" · risposta aperta inclusa" if open_answers else ""), deeplink=f"/courses/{course_id}", dedupe_key=f"course_test_{course_id}_{user['user_id']}", team_copy=False)
    except Exception:  # noqa: BLE001
        pass
    await _course_cert_issue_if_ready(c, user["user_id"])  # V50.z276
    return {"ok": True, "score": score, "total": total}


# --- Relazione vocale fine corso: trascrizione (VoiceDictationButton → Whisper lato app) → Axon organizza → editor inline ---
COURSE_REPORT_SYSTEM = (
    "Sei AxonOS. Ricevi la trascrizione (parlato, in italiano) della relazione di fine corso registrata dal formatore. "
    "Organizzala e restituisci SOLO JSON: {\"summary\": \"riassunto 5-8 righe\", \"highlights\": [\"punto saliente\", ...], "
    "\"actions\": [\"azione concreta / follow-up\", ...], \"feedback\": \"feedback ricevuti o percepiti sui partecipanti\", \"proposals\": [\"proposta di miglioramento\", ...]}. "
    "Non inventare: se un elemento non è presente lascia lista vuota o stringa vuota. Italiano, tono professionale."
)


@api_router.post("/courses/{course_id}/report/organize")
async def course_report_organize(course_id: str, body: CourseReportIn, user=Depends(get_current_user)):
    c = await _course_creator_only(course_id, user)
    tr = (body.transcript or "").strip()
    if len(tr) < 10:
        raise HTTPException(status_code=400, detail="Trascrizione troppo breve")
    try:
        data = await _ft_llm(COURSE_REPORT_SYSTEM, f"Corso: {c['title']}\n\nTRASCRIZIONE:\n{tr[:12000]}", session="course-report")
    except HTTPException:
        data = {"summary": tr[:1200], "highlights": [], "actions": [], "feedback": "", "proposals": []}
    rep = {"transcript": tr[:20000], "summary": str(data.get("summary") or "")[:4000], "highlights": [str(x)[:300] for x in (data.get("highlights") or []) if str(x).strip()][:12],
           "actions": [str(x)[:300] for x in (data.get("actions") or []) if str(x).strip()][:12], "feedback": str(data.get("feedback") or "")[:2000],
           "proposals": [str(x)[:300] for x in (data.get("proposals") or []) if str(x).strip()][:12], "visible_to_recipients": bool((c.get("report") or {}).get("visible_to_recipients", False)),
           "author": user.get("name") or user.get("email"), "updated_at": utcnow()}
    await db.courses_folders.update_one({"course_id": course_id}, {"$set": {"report": rep}})
    return {"report": rep}


@api_router.put("/courses/{course_id}/report")
async def course_report_update(course_id: str, body: CourseReportIn, user=Depends(get_current_user)):
    c = await _course_creator_only(course_id, user)
    rep = dict(c.get("report") or {})
    for k in ("transcript", "summary", "highlights", "actions", "feedback", "proposals", "visible_to_recipients"):
        v = getattr(body, k)
        if v is not None:
            rep[k] = v
    rep["updated_at"] = utcnow()
    await db.courses_folders.update_one({"course_id": course_id}, {"$set": {"report": rep}})
    return {"report": rep}


@api_router.get("/courses/{course_id}/report")
async def course_report_get(course_id: str, user=Depends(get_current_user)):
    c = await _course_get(course_id, user)
    rep = c.get("report")
    creator = c.get("created_by") == user["user_id"] or is_owner(user)
    if not rep or (not creator and not rep.get("visible_to_recipients")):
        return {"report": None, "can_edit": creator}
    if not creator:
        rep = {k: v for k, v in rep.items() if k != "transcript"}
    return {"report": rep, "can_edit": creator}


# --- Dashboard azienda (Master): tutti i corsi del tenant con filtri manager / area(gruppo) / periodo ---
@api_router.get("/courses-dashboard")  # (non /courses/dashboard: collide con /courses/{course_id})
async def courses_dashboard(manager_id: Optional[str] = None, area: Optional[str] = None, date_from: Optional[str] = None, date_to: Optional[str] = None,
                            include_archived: bool = True, user=Depends(get_current_user)):
    if not await _course_can_manage(user):
        raise HTTPException(status_code=403, detail="Solo Master/Manager")
    q: Dict[str, Any] = {"tenant_id": tenant_id(user)}
    if not is_owner(user):  # Manager: solo i propri corsi
        q["created_by"] = user["user_id"]
    if manager_id:
        q["created_by"] = manager_id
    if not include_archived:
        q["status"] = "active"
    if date_from or date_to:
        rng: Dict[str, Any] = {}
        if date_from:
            rng["$gte"] = date_from[:10]
        if date_to:
            rng["$lte"] = date_to[:10]
        q["course_date"] = rng
    courses = await db.courses_folders.find(q, {"_id": 0, "files": 0, "body_text": 0, "report": 0}).sort("course_date", -1).to_list(500)
    # area/gruppo = campo `area` o `group` dell'utente (se presente), altrimenti il Manager di riferimento
    all_ids = {c["created_by"] for c in courses} | {r for c in courses for r in (c.get("recipients") or [])}
    users: Dict[str, dict] = {}
    async for u in db.users.find({"user_id": {"$in": list(all_ids)}}, {"_id": 0, "user_id": 1, "name": 1, "email": 1, "area": 1, "group": 1, "team_name": 1, "parent_user_id": 1}):
        users[u["user_id"]] = u
    def area_of(uid: str) -> str:
        u = users.get(uid) or {}
        return str(u.get("area") or u.get("group") or u.get("team_name") or (users.get(u.get("parent_user_id") or "") or {}).get("name") or "Senza area")
    tests = {t["course_id"]: t for t in await db.courses_tests.find({"course_id": {"$in": [c["course_id"] for c in courses]}}, {"_id": 0, "course_id": 1, "responses": 1}).to_list(500)}
    rows, by_manager, by_area = [], {}, {}
    def acc(bucket: dict, key: str, label: str, row: dict):
        b = bucket.setdefault(key, {"key": key, "label": label, "courses": 0, "recipients": 0, "reads": 0, "yes": 0, "no": 0, "tests": 0, "score_sum": 0.0, "score_n": 0})
        b["courses"] += 1; b["recipients"] += row["recipients"]; b["reads"] += row["reads"]; b["yes"] += row["yes"]; b["no"] += row["no"]; b["tests"] += row["test_responses"]
        if row["avg_score_pct"] is not None:
            b["score_sum"] += row["avg_score_pct"]; b["score_n"] += 1
    for c in courses:
        rec = c.get("recipients") or []
        if area and not any(area_of(r) == area for r in rec) and area_of(c["created_by"]) != area:
            continue
        reads, part = c.get("reads") or {}, c.get("participation") or {}
        t = tests.get(c["course_id"]) or {}
        resp = t.get("responses") or {}
        scored = [r for r in resp.values() if r.get("total")]
        row = {"course_id": c["course_id"], "title": c["title"], "course_date": c.get("course_date") or "", "status": c.get("status"), "manager_id": c["created_by"],
               "manager": (users.get(c["created_by"]) or {}).get("name") or "", "recipients": len(rec), "reads": sum(1 for r in rec if r in reads),
               "yes": sum(1 for r in rec if (part.get(r) or {}).get("status") == "yes"), "no": sum(1 for r in rec if (part.get(r) or {}).get("status") == "no"),
               "test_responses": len(resp), "avg_score_pct": round(sum(r["score"] / r["total"] * 100 for r in scored) / len(scored), 1) if scored else None,
               "areas": sorted({area_of(r) for r in rec})}
        row["participation_pct"] = round(row["yes"] / row["recipients"] * 100) if row["recipients"] else 0
        rows.append(row)
        acc(by_manager, c["created_by"], row["manager"] or c["created_by"], row)
        for a in row["areas"] or ["Senza area"]:
            acc(by_area, a, a, row)
    def fin(b: dict) -> dict:
        b["participation_pct"] = round(b["yes"] / b["recipients"] * 100) if b["recipients"] else 0
        b["read_pct"] = round(b["reads"] / b["recipients"] * 100) if b["recipients"] else 0
        b["avg_score_pct"] = round(b["score_sum"] / b["score_n"], 1) if b["score_n"] else None
        b.pop("score_sum"); b.pop("score_n"); return b
    managers = [{"user_id": uid, "name": (users.get(uid) or {}).get("name") or uid} for uid in sorted({c["created_by"] for c in courses})]
    areas = sorted({area_of(r) for c in courses for r in (c.get("recipients") or [])})
    tot = {"courses": len(rows), "recipients": sum(r["recipients"] for r in rows), "reads": sum(r["reads"] for r in rows), "yes": sum(r["yes"] for r in rows), "no": sum(r["no"] for r in rows), "tests": sum(r["test_responses"] for r in rows)}
    tot["participation_pct"] = round(tot["yes"] / tot["recipients"] * 100) if tot["recipients"] else 0
    return {"courses": rows, "totals": tot, "by_manager": [fin(b) for b in by_manager.values()], "by_area": [fin(b) for b in by_area.values()], "filters": {"managers": managers, "areas": areas}}


# =====================================================================================
# V50.z275 — Fase 3: GENERAZIONE POWERPOINT (AI dai materiali + scrittura manuale del Master)
# Bozza slide salvata nel corso (`slides`), editor inline, export .pptx (python-pptx), opzione "aggiungi ai materiali".
# =====================================================================================
class SlideIn(BaseModel):
    title: str = ""
    bullets: List[str] = Field(default_factory=list)
    notes: Optional[str] = ""


class SlidesIn(BaseModel):
    title: Optional[str] = None
    subtitle: Optional[str] = None
    slides: Optional[List[SlideIn]] = None
    visible_to_recipients: Optional[bool] = None


class SlidesGenerateIn(BaseModel):
    n_slides: int = 8
    instructions: Optional[str] = ""


COURSE_SLIDES_SYSTEM = (
    "Sei AxonOS, formatore aziendale per consulenti energetici. Ricevi i materiali di un corso (testo, estratti PDF, eventuali immagini). "
    "Crea una presentazione PowerPoint chiara e didattica in italiano. Restituisci SOLO JSON: "
    "{\"title\": \"titolo presentazione\", \"subtitle\": \"sottotitolo breve\", \"slides\": [{\"title\": \"titolo slide (max 8 parole)\", "
    "\"bullets\": [\"punto chiave conciso (max 18 parole)\", ...], \"notes\": \"note per il relatore, 2-4 frasi\"}, ...]}. "
    "Ogni slide 3-5 bullet. Prima slide = agenda/obiettivi, ultima = riepilogo/conclusioni. Non inventare dati non presenti nei materiali. "
    "Se il formatore ha dettato indicazioni (struttura, argomenti, tono, esempi), seguile fedelmente: se i materiali sono scarsi, costruisci le slide a partire da ciò che ha dettato."
)

_OFFICE_XML_TAG = re.compile(r"<[^>]+>")


def _course_office_text(raw: bytes, kind: str) -> str:
    """Estrae testo da docx/pptx/xlsx leggendo direttamente gli XML (zero dipendenze extra)."""
    import zipfile
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            names = sorted(n for n in z.namelist() if (
                (kind == "docx" and n == "word/document.xml") or (kind == "pptx" and re.match(r"ppt/slides/slide\d+\.xml", n))
                or (kind == "xlsx" and n == "xl/sharedStrings.xml")))
            out = []
            for n in names[:60]:
                xml = z.read(n).decode("utf-8", "ignore")
                xml = re.sub(r"</(a:p|w:p|si)>", "\n", xml)
                out.append(html_lib.unescape(_OFFICE_XML_TAG.sub(" ", xml)))
            return re.sub(r"[ \t]+", " ", "\n".join(out)).strip()
    except Exception:  # noqa: BLE001
        return ""


async def _course_file_bytes(f: dict) -> Optional[bytes]:
    if f.get("path"):
        got = await run_in_threadpool(_get_object_sync, f["path"])
        if got:
            return got[0]
    if f.get("data_b64"):
        try:
            return base64.b64decode(f["data_b64"])
        except Exception:  # noqa: BLE001
            return None
    return None


async def _course_materials(c: dict) -> Tuple[str, List[str]]:
    """Testo + immagini (b64 png/jpg) dai materiali del corso, per il prompt AI."""
    parts: List[str] = [f"TITOLO CORSO: {c.get('title')}"]
    if c.get("description"):
        parts.append(f"DESCRIZIONE: {c['description']}")
    if c.get("body_text"):
        parts.append(f"APPUNTI DEL FORMATORE:\n{c['body_text'][:12000]}")
    images: List[str] = []
    for f in (c.get("files") or [])[:12]:
        name, mime = (f.get("name") or "").lower(), (f.get("mime") or "").lower()
        raw = await _course_file_bytes(f)
        if not raw:
            continue
        txt = ""
        if mime == "application/pdf" or name.endswith(".pdf"):
            try:
                from pypdf import PdfReader
                reader = PdfReader(io.BytesIO(raw))
                txt = "\n".join((p.extract_text() or "") for p in reader.pages[:25])
            except Exception:  # noqa: BLE001
                txt = ""
            if len(txt.strip()) < 200 and len(images) < 6:  # PDF scansionato → pagine come immagini
                try:
                    for p in _ft_render_pages(raw, dpi=70)[:3]:
                        images.append(p["png_b64"])
                except Exception:  # noqa: BLE001
                    pass
        elif mime.startswith("image/"):
            if len(images) < 6 and len(raw) < 4 * 1024 * 1024:
                images.append(base64.b64encode(raw).decode())
        elif name.endswith(".docx"):
            txt = _course_office_text(raw, "docx")
        elif name.endswith(".pptx"):
            txt = _course_office_text(raw, "pptx")
        elif name.endswith(".xlsx"):
            txt = _course_office_text(raw, "xlsx")
        elif mime.startswith("text/") or name.endswith((".txt", ".md", ".csv")):
            txt = raw.decode("utf-8", "ignore")
        if txt.strip():
            parts.append(f"--- MATERIALE «{f.get('name')}» ---\n{txt.strip()[:15000]}")
    return "\n\n".join(parts)[:60000], images


def _course_slides_norm(title: str, subtitle: str, slides: List[dict]) -> List[dict]:
    out = []
    for sl in slides[:40]:
        t = str(sl.get("title") or "").strip()[:120]
        bl = [str(b).strip()[:300] for b in (sl.get("bullets") or []) if str(b).strip()][:8]
        if not t and not bl:
            continue
        out.append({"title": t or "Slide", "bullets": bl, "notes": str(sl.get("notes") or "").strip()[:2000]})
    return out


def _course_slides_public(c: dict, creator: bool) -> dict:
    sl = c.get("slides")
    if not sl or (not creator and not sl.get("visible_to_recipients")):
        return {"slides": None, "can_edit": creator}
    return {"slides": {k: v for k, v in sl.items()}, "can_edit": creator}


@api_router.post("/courses/{course_id}/slides/generate")
async def course_slides_generate(course_id: str, body: SlidesGenerateIn, user=Depends(get_current_user)):
    """AI: legge appunti + allegati (PDF/immagini/office/testo) → bozza slide (salvata, modificabile)."""
    c = await _course_creator_only(course_id, user)
    text, images = await _course_materials(c)
    instr = (body.instructions or "").strip()[:4000]
    if len(text) < 60 and not images and len(instr) < 40:
        raise HTTPException(status_code=400, detail="Materiali insufficienti: aggiungi appunti o allegati al corso, oppure detta a voce cosa vuoi nelle slide")
    n = max(3, min(20, int(body.n_slides or 8)))
    prompt = (f"Genera esattamente {n} slide.\n" + (f"INDICAZIONI DEL FORMATORE (dettate a voce o scritte — hanno priorità sui materiali e vanno seguite fedelmente):\n{instr}\n" if instr else "")
              + f"\nMATERIALI DEL CORSO:\n{text}")
    data = await _ft_llm(COURSE_SLIDES_SYSTEM, prompt, images or None, session="course-slides")
    slides = _course_slides_norm(str(data.get("title") or ""), str(data.get("subtitle") or ""), data.get("slides") or [])
    if not slides:
        raise HTTPException(status_code=422, detail="L'AI non ha prodotto slide valide: riprova o scrivile manualmente")
    prev = c.get("slides") or {}
    doc = {"title": str(data.get("title") or c["title"]).strip()[:160] or c["title"], "subtitle": str(data.get("subtitle") or "").strip()[:200], "slides": slides,
           "source": "ai", "visible_to_recipients": bool(prev.get("visible_to_recipients", False)), "author": user.get("name") or user.get("email"), "updated_at": utcnow()}
    await db.courses_folders.update_one({"course_id": course_id}, {"$set": {"slides": doc}})
    return {"slides": doc, "can_edit": True, "materials_chars": len(text), "images": len(images)}


@api_router.put("/courses/{course_id}/slides")
async def course_slides_update(course_id: str, body: SlidesIn, user=Depends(get_current_user)):
    """Master scrive/modifica manualmente le slide (titoli, punti, note relatore)."""
    c = await _course_creator_only(course_id, user)
    doc = dict(c.get("slides") or {"title": c["title"], "subtitle": "", "slides": [], "source": "manual", "visible_to_recipients": False})
    if body.title is not None:
        doc["title"] = body.title.strip()[:160] or c["title"]
    if body.subtitle is not None:
        doc["subtitle"] = body.subtitle.strip()[:200]
    if body.slides is not None:
        doc["slides"] = _course_slides_norm(doc["title"], doc.get("subtitle") or "", [s.model_dump() for s in body.slides])
        doc["source"] = "manual" if doc.get("source") != "ai" else "ai+manual"
    if body.visible_to_recipients is not None:
        doc["visible_to_recipients"] = bool(body.visible_to_recipients)
    doc["author"] = user.get("name") or user.get("email")
    doc["updated_at"] = utcnow()
    await db.courses_folders.update_one({"course_id": course_id}, {"$set": {"slides": doc}})
    return {"slides": doc, "can_edit": True}


@api_router.delete("/courses/{course_id}/slides")
async def course_slides_delete(course_id: str, user=Depends(get_current_user)):
    await _course_creator_only(course_id, user)
    await db.courses_folders.update_one({"course_id": course_id}, {"$unset": {"slides": ""}})
    return {"ok": True}


@api_router.get("/courses/{course_id}/slides")
async def course_slides_get(course_id: str, user=Depends(get_current_user)):
    c = await _course_get(course_id, user)
    return _course_slides_public(c, c.get("created_by") == user["user_id"] or is_owner(user))


def _course_build_pptx(c: dict, sl: dict) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    PURPLE, DARK, MUTED, WHITE = RGBColor(0x6B, 0x21, 0xA8), RGBColor(0x1E, 0x1B, 0x4B), RGBColor(0x64, 0x74, 0x8B), RGBColor(0xFF, 0xFF, 0xFF)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    W, H = prs.slide_width, prs.slide_height
    date_txt = (c.get("course_date") or "").split("-")
    date_txt = "/".join(reversed(date_txt)) if len(date_txt) == 3 else ""
    footer = " · ".join(x for x in (c.get("title") or "", date_txt, c.get("location") or "") if x)

    def add_text(slide, left, top, width, height, text, size, bold=False, color=DARK, align=None):
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size, p.font.bold, p.font.color.rgb = Pt(size), bold, color
        if align is not None:
            p.alignment = align
        return tb

    # --- Copertina ---
    s0 = prs.slides.add_slide(blank)
    bg = s0.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid(); bg.fill.fore_color.rgb = PURPLE; bg.line.fill.background()
    add_text(s0, Inches(0.9), Inches(2.4), W - Inches(1.8), Inches(1.6), sl.get("title") or c["title"], 40, True, WHITE)
    if sl.get("subtitle"):
        add_text(s0, Inches(0.9), Inches(4.0), W - Inches(1.8), Inches(1.0), sl["subtitle"], 20, False, WHITE)
    add_text(s0, Inches(0.9), Inches(6.4), W - Inches(1.8), Inches(0.5), footer, 14, False, WHITE)

    # --- Contenuti ---
    for i, s in enumerate(sl.get("slides") or [], start=1):
        sd = prs.slides.add_slide(blank)
        bar = sd.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(1.25))
        bar.fill.solid(); bar.fill.fore_color.rgb = PURPLE; bar.line.fill.background()
        add_text(sd, Inches(0.7), Inches(0.25), W - Inches(1.4), Inches(0.9), s.get("title") or f"Slide {i}", 28, True, WHITE)
        tb = sd.shapes.add_textbox(Inches(0.9), Inches(1.7), W - Inches(1.8), Inches(4.8))
        tf = tb.text_frame; tf.word_wrap = True
        for k, b in enumerate(s.get("bullets") or []):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.text = f"•  {b}"
            p.font.size, p.font.color.rgb = Pt(20 if len(s.get("bullets") or []) <= 5 else 17), DARK
            p.space_after = Pt(12)
        add_text(sd, Inches(0.9), Inches(6.9), W - Inches(3.0), Inches(0.4), footer, 11, False, MUTED)
        add_text(sd, W - Inches(2.0), Inches(6.9), Inches(1.3), Inches(0.4), f"{i} / {len(sl.get('slides') or [])}", 11, False, MUTED)
        if s.get("notes"):
            sd.notes_slide.notes_text_frame.text = s["notes"]

    # --- Chiusura ---
    s9 = prs.slides.add_slide(blank)
    bg = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid(); bg.fill.fore_color.rgb = PURPLE; bg.line.fill.background()
    add_text(s9, Inches(0.9), Inches(3.0), W - Inches(1.8), Inches(1.2), "Grazie per l'attenzione", 36, True, WHITE)
    add_text(s9, Inches(0.9), Inches(4.2), W - Inches(1.8), Inches(0.6), f"Relatore: {sl.get('author') or ''}".strip(), 16, False, WHITE)
    buf = io.BytesIO(); prs.save(buf)
    return buf.getvalue()


PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


@api_router.get("/courses/{course_id}/slides/pptx")
async def course_slides_pptx(course_id: str, user=Depends(get_current_user)):
    """Scarica la presentazione .pptx (creatore sempre; collaboratori se resa visibile)."""
    c = await _course_get(course_id, user)
    creator = c.get("created_by") == user["user_id"] or is_owner(user)
    sl = c.get("slides")
    if not sl or not sl.get("slides") or (not creator and not sl.get("visible_to_recipients")):
        raise HTTPException(status_code=404, detail="Nessuna presentazione disponibile")
    raw = await run_in_threadpool(_course_build_pptx, c, sl)
    safe = re.sub(r"[^A-Za-z0-9_\- ]+", "", sl.get("title") or c["title"]).strip() or "presentazione"
    return {"name": f"{safe}.pptx", "mime": PPTX_MIME, "data_base64": base64.b64encode(raw).decode()}


@api_router.post("/courses/{course_id}/slides/attach")
async def course_slides_attach(course_id: str, user=Depends(get_current_user)):
    """Genera il .pptx e lo aggiunge ai materiali del corso (sostituisce la versione precedente)."""
    c = await _course_creator_only(course_id, user)
    sl = c.get("slides")
    if not sl or not sl.get("slides"):
        raise HTTPException(status_code=400, detail="Crea prima le slide")
    raw = await run_in_threadpool(_course_build_pptx, c, sl)
    safe = re.sub(r"[^A-Za-z0-9_\- ]+", "", sl.get("title") or c["title"]).strip() or "presentazione"
    files = [f for f in (c.get("files") or []) if not f.get("from_slides")]
    meta = await _course_store_file(c["tenant_id"], course_id, CourseFileIn(name=f"{safe}.pptx", mime=PPTX_MIME, data_base64=base64.b64encode(raw).decode()))
    meta["from_slides"] = True
    files.append(meta)
    await db.courses_folders.update_one({"course_id": course_id}, {"$set": {"files": files, "updated_at": utcnow()}})
    return {"ok": True, "file": {k: v for k, v in meta.items() if k != "data_b64"}}


# =====================================================================================
# V50.z276 — ATTESTATO DI PARTECIPAZIONE (PDF automatico) · SLIDE DA VOCE (dettatura → istruzioni AI)
# Idoneo = ha risposto ✅ Partecipo E (nessun test oppure test completato). Quando diventa idoneo riceve la notifica
# «🎓 Attestato disponibile» (deeplink al corso) e il PDF è scaricabile da lui e dal creatore.
# =====================================================================================
def _course_cert_eligible(c: dict, t: Optional[dict], uid: str) -> bool:
    if (c.get("certificates") or {}).get(uid):  # V50.z300 — già emesso (anche dal Master con «Rilascia a tutti»)
        return True
    if (c.get("participation") or {}).get(uid, {}).get("status") != "yes":
        return False
    # V50.z302 — PRESENZE: se il formatore ha registrato le presenze, l'attestato spetta SOLO ai presenti;
    # se il corso ha una data futura, l'attestato non è ancora disponibile (si rilascia dopo la lezione).
    att = c.get("attendance") or {}
    if c.get("attendance_taken_at"):
        if not (att.get(uid) or {}).get("present"):
            return False
    elif c.get("course_date") and c["course_date"] > _course_today_str():
        return False
    if t and t.get("questions"):
        return uid in (t.get("responses") or {})
    return True


async def _course_cert_issue_if_ready(c: dict, uid: str) -> None:
    """Chiamata dopo partecipazione/test: se idoneo e non ancora emesso → registra + notifica."""
    c = await db.courses_folders.find_one({"course_id": c["course_id"]}, {"_id": 0, "files": 0, "body_text": 0}) or c
    if (c.get("certificates") or {}).get(uid):
        return
    t = await db.courses_tests.find_one({"course_id": c["course_id"]}, {"_id": 0, "questions": 1, "responses": 1})
    if not _course_cert_eligible(c, t, uid):
        return
    cert = {"cert_id": f"ATT-{utcnow().strftime('%Y%m%d')}-{uuid.uuid4().hex[:6].upper()}", "issued_at": utcnow()}
    await db.courses_folders.update_one({"course_id": c["course_id"]}, {"$set": {f"certificates.{uid}": cert}})
    try:
        await axon_notify(uid, c["tenant_id"], "reminder", f"🎓 Attestato disponibile: {c['title']}", "Hai completato il corso: scarica il tuo attestato di partecipazione in PDF dalla cartella del corso.",
                          deeplink=f"/courses/{c['course_id']}", dedupe_key=f"course_cert_{c['course_id']}_{uid}", team_copy=False)
    except Exception:  # noqa: BLE001
        pass


def _course_build_certificate_pdf(c: dict, cert: dict, person: str, trainer: str, company: str, logo_b64: str, score: Optional[Tuple[int, int]]) -> bytes:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfgen import canvas as _canvas
    from reportlab.lib.colors import HexColor
    from reportlab.lib.utils import ImageReader
    W, H = landscape(A4)
    buf = io.BytesIO()
    cv = _canvas.Canvas(buf, pagesize=(W, H))
    purple, gold, dark, muted = HexColor("#6B21A8"), HexColor("#D97706"), HexColor("#1E1B4B"), HexColor("#64748B")
    cv.setFillColor(HexColor("#FAF5FF")); cv.rect(0, 0, W, H, stroke=0, fill=1)
    cv.setStrokeColor(purple); cv.setLineWidth(6); cv.rect(28, 28, W - 56, H - 56, stroke=1, fill=0)
    cv.setStrokeColor(gold); cv.setLineWidth(1.5); cv.rect(40, 40, W - 80, H - 80, stroke=1, fill=0)
    y = H - 90
    if logo_b64:
        try:
            raw = base64.b64decode(logo_b64.split(",", 1)[1] if logo_b64.startswith("data:") else logo_b64)
            img = ImageReader(io.BytesIO(raw)); iw, ih = img.getSize(); h = 54; w = iw * h / ih
            cv.drawImage(img, (W - w) / 2, y - 30, w, h, mask="auto"); y -= 70
        except Exception:  # noqa: BLE001
            pass
    if company:
        cv.setFillColor(muted); cv.setFont("Helvetica-Bold", 13); cv.drawCentredString(W / 2, y, company.upper()); y -= 40
    cv.setFillColor(purple); cv.setFont("Helvetica-Bold", 34); cv.drawCentredString(W / 2, y, "ATTESTATO DI PARTECIPAZIONE"); y -= 42
    cv.setFillColor(muted); cv.setFont("Helvetica", 14); cv.drawCentredString(W / 2, y, "Si certifica che"); y -= 44
    cv.setFillColor(dark); cv.setFont("Helvetica-Bold", 30); cv.drawCentredString(W / 2, y, person); y -= 18
    cv.setStrokeColor(gold); cv.setLineWidth(1); cv.line(W / 2 - 200, y, W / 2 + 200, y); y -= 34
    cv.setFillColor(muted); cv.setFont("Helvetica", 14); cv.drawCentredString(W / 2, y, "ha partecipato con esito positivo al corso di formazione"); y -= 36
    cv.setFillColor(purple); cv.setFont("Helvetica-Bold", 22)
    title = c.get("title") or "Corso"
    cv.drawCentredString(W / 2, y, title[:70] + ("…" if len(title) > 70 else "")); y -= 32
    d = (c.get("course_date") or "").split("-")
    info = " · ".join(x for x in ("/".join(reversed(d)) if len(d) == 3 else "", c.get("location") or "") if x)
    cv.setFillColor(dark); cv.setFont("Helvetica", 13)
    if info:
        cv.drawCentredString(W / 2, y, info); y -= 22
    if score and score[1]:
        cv.drawCentredString(W / 2, y, f"Test finale superato: {score[0]}/{score[1]} risposte corrette ({round(score[0] / score[1] * 100)}%)"); y -= 22
    # firme
    yb = 95
    cv.setStrokeColor(dark); cv.setLineWidth(0.8)
    cv.line(110, yb, 330, yb); cv.line(W - 330, yb, W - 110, yb)
    cv.setFillColor(dark); cv.setFont("Helvetica-Bold", 11); cv.drawCentredString(220, yb - 16, trainer or "Il formatore"); cv.drawCentredString(W - 220, yb - 16, company or "La direzione")
    cv.setFillColor(muted); cv.setFont("Helvetica", 9); cv.drawCentredString(220, yb - 29, "Formatore"); cv.drawCentredString(W - 220, yb - 29, "Azienda")
    issued = cert.get("issued_at")
    issued_txt = issued.astimezone().strftime("%d/%m/%Y") if isinstance(issued, datetime) else ""
    cv.setFont("Helvetica", 9); cv.drawCentredString(W / 2, 52, f"Attestato n. {cert.get('cert_id')} · rilasciato il {issued_txt} · generato da AdvisorOS")
    cv.showPage(); cv.save()
    return buf.getvalue()


@api_router.get("/courses/{course_id}/certificate")
async def course_certificate(course_id: str, user_id: Optional[str] = None, user=Depends(get_current_user)):
    """PDF attestato: il collaboratore scarica il proprio; il creatore può scaricare quello di ogni partecipante (?user_id=)."""
    c = await _course_get(course_id, user)
    creator = c.get("created_by") == user["user_id"] or is_owner(user)
    uid = user_id if (user_id and creator) else user["user_id"]
    t = await db.courses_tests.find_one({"course_id": course_id}, {"_id": 0, "questions": 1, "responses": 1})
    if not _course_cert_eligible(c, t, uid):
        raise HTTPException(status_code=400, detail="Attestato non ancora disponibile: serve ✅ Partecipo e il test completato")
    cert = (c.get("certificates") or {}).get(uid)
    if not cert:
        cert = {"cert_id": f"ATT-{utcnow().strftime('%Y%m%d')}-{uuid.uuid4().hex[:6].upper()}", "issued_at": utcnow()}
        await db.courses_folders.update_one({"course_id": course_id}, {"$set": {f"certificates.{uid}": cert}})
    names = await _course_names([uid, c["created_by"]])
    owner = await db.users.find_one({"user_id": c["tenant_id"]}, {"_id": 0, "company_name": 1, "company_logo_base64": 1}) or {}
    resp = ((t or {}).get("responses") or {}).get(uid) or {}
    score = (int(resp.get("score") or 0), int(resp.get("total") or 0)) if resp else None
    raw = await run_in_threadpool(_course_build_certificate_pdf, c, cert, names.get(uid, uid), names.get(c["created_by"], ""), (owner.get("company_name") or "").strip(), owner.get("company_logo_base64") or "", score)
    safe = re.sub(r"[^A-Za-z0-9_\- ]+", "", f"Attestato {names.get(uid, '')} {c['title']}").strip() or "attestato"
    return {"name": f"{safe}.pdf", "mime": "application/pdf", "data_base64": base64.b64encode(raw).decode(), "cert": cert}
