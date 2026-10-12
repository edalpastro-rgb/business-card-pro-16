# V50.z536 - Backend tests for Knowledge Base multi-format, OG preview, FAQ suggest
import base64
import io
import os
import time
from typing import List

import pytest
import requests

BASE_URL = "http://localhost:8001"
API = f"{BASE_URL}/api"

OWNER_EMAIL = "test@dalpastro.it"
OWNER_PASSWORD = "Test1234!"
CMS_PASSWORD = (os.environ.get("LANDING_CMS_PASSWORD") or "axon-cms-2026")

FONT_PATH = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"

CREATED_DOCS: List[str] = []


# ---------- Fixtures ----------

@pytest.fixture(scope="module")
def auth_token():
    r = requests.post(f"{API}/auth/login", json={"email": OWNER_EMAIL, "password": OWNER_PASSWORD}, timeout=15)
    assert r.status_code == 200, f"login failed: {r.status_code} {r.text}"
    tok = r.json().get("token")
    assert tok
    return tok


@pytest.fixture(scope="module")
def auth_headers(auth_token):
    return {"Authorization": f"Bearer {auth_token}", "Content-Type": "application/json"}


# ---------- Helpers ----------

def _poll_doc(headers, doc_id, timeout=120, interval=2):
    """Poll kb_documents until status != indexing. Returns final doc dict or None."""
    deadline = time.time() + timeout
    last = None
    while time.time() < deadline:
        r = requests.get(f"{API}/kb/documents", headers=headers, timeout=15)
        if r.status_code != 200:
            time.sleep(interval); continue
        items = r.json().get("items", [])
        for it in items:
            if it.get("doc_id") == doc_id:
                last = it
                if it.get("status") in ("ready", "error"):
                    return it
        time.sleep(interval)
    return last


def _delete_doc(headers, doc_id):
    try:
        requests.delete(f"{API}/kb/documents/{doc_id}", headers=headers, timeout=10)
    except Exception:
        pass


def _b64(data: bytes) -> str:
    return base64.b64encode(data).decode()


def _gen_docx(text: str) -> bytes:
    import docx
    buf = io.BytesIO()
    d = docx.Document()
    for line in text.split("\n"):
        d.add_paragraph(line)
    d.save(buf)
    return buf.getvalue()


def _gen_xlsx(rows) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _gen_png_with_text(text: str, w=1200, h=400) -> bytes:
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (w, h), (255, 255, 255))
    d = ImageDraw.Draw(img)
    font = ImageFont.truetype(FONT_PATH, 48)
    d.text((40, 120), text, fill=(0, 0, 0), font=font)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _gen_blank_png() -> bytes:
    from PIL import Image
    img = Image.new("RGB", (800, 400), (255, 255, 255))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _gen_image_pdf(text: str) -> bytes:
    """PDF composed of image pages (no text layer) using pymupdf insert_image."""
    import pymupdf as fitz
    img_bytes = _gen_png_with_text(text, w=1200, h=800)
    doc = fitz.open()
    for _ in range(2):
        page = doc.new_page(width=1200, height=800)
        page.insert_image(fitz.Rect(0, 0, 1200, 800), stream=img_bytes)
    out = doc.tobytes()
    doc.close()
    return out


def _gen_text_pdf(text: str) -> bytes:
    """PDF with real text layer."""
    import pymupdf as fitz
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text, fontsize=14)
    out = doc.tobytes()
    doc.close()
    return out


# ---------- 1. KB status ----------

class TestKbStatus:
    def test_status(self, auth_headers):
        r = requests.get(f"{API}/kb/status", headers=auth_headers, timeout=10)
        assert r.status_code == 200, r.text
        data = r.json()
        assert data.get("ocr_ready") is True, f"ocr_ready expected True, got {data.get('ocr_ready')}"
        fmts = data.get("formats") or []
        for f in ("docx", "xlsx", "png"):
            assert f in fmts, f"format {f} missing from {fmts}"
        assert data.get("embedder_ready") is True


# ---------- 2. DOCX ----------

class TestKbDocx:
    def test_docx_upload(self, auth_headers):
        txt = "Questo è un documento di test DOCX. Contiene informazioni sui crediti AI e sul sistema di knowledge base. " * 3
        raw = _gen_docx(txt)
        payload = {"title": "TEST_docx", "filename": "test.docx", "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "data_base64": _b64(raw)}
        r = requests.post(f"{API}/kb/documents", headers=auth_headers, json=payload, timeout=20)
        assert r.status_code == 200, f"status={r.status_code} body={r.text}"
        body = r.json()
        doc = body["document"]
        assert doc["source_type"] == "docx"
        assert body["ocr"] is False
        doc_id = doc["doc_id"]
        CREATED_DOCS.append(doc_id)
        final = _poll_doc(auth_headers, doc_id, timeout=60)
        assert final is not None, "doc disappeared"
        assert final["status"] == "ready", f"status={final.get('status')} err={final.get('error')}"
        assert final.get("chunks", 0) >= 1


# ---------- 3. XLSX ----------

class TestKbXlsx:
    def test_xlsx_upload(self, auth_headers):
        rows = [["Nome", "Categoria", "Prezzo"], ["Mela", "Frutta", 1.2], ["Pane", "Panetteria", 2.5], ["Formaggio", "Latticini", 8.9]]
        raw = _gen_xlsx(rows)
        payload = {"title": "TEST_xlsx", "filename": "test.xlsx", "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "data_base64": _b64(raw)}
        r = requests.post(f"{API}/kb/documents", headers=auth_headers, json=payload, timeout=20)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["document"]["source_type"] == "xlsx"
        doc_id = body["document"]["doc_id"]
        CREATED_DOCS.append(doc_id)
        final = _poll_doc(auth_headers, doc_id, timeout=60)
        assert final["status"] == "ready", f"err={final.get('error')}"
        assert final.get("chunks", 0) >= 1


# ---------- 4. PNG with text (OCR) + search ----------

class TestKbPngOcr:
    def test_png_ocr_and_search(self, auth_headers):
        secret = "La capitale della Francia e' Parigi e il fiume principale e' la Senna"
        raw = _gen_png_with_text(secret)
        payload = {"title": "TEST_png_ocr", "filename": "test.png", "mime": "image/png", "data_base64": _b64(raw)}
        r = requests.post(f"{API}/kb/documents", headers=auth_headers, json=payload, timeout=20)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["document"]["source_type"] == "image"
        assert body["ocr"] is True
        doc_id = body["document"]["doc_id"]
        CREATED_DOCS.append(doc_id)
        final = _poll_doc(auth_headers, doc_id, timeout=120)
        assert final["status"] == "ready", f"err={final.get('error')}"
        # Search
        sr = requests.post(f"{API}/kb/search", headers=auth_headers, json={"query": "qual e' la capitale della Francia", "k": 3}, timeout=15)
        assert sr.status_code == 200
        hits = sr.json().get("hits", [])
        assert any(h["doc_id"] == doc_id and h["similarity"] >= 0.38 for h in hits), f"no good hit: {hits}"


# ---------- 5. Image-only PDF ----------

class TestKbScannedPdf:
    def test_image_pdf_ocr(self, auth_headers):
        raw = _gen_image_pdf("Contratto di prova con clausole importanti")
        payload = {"title": "TEST_scanned_pdf", "filename": "scan.pdf", "mime": "application/pdf", "data_base64": _b64(raw)}
        r = requests.post(f"{API}/kb/documents", headers=auth_headers, json=payload, timeout=30)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["document"]["source_type"] == "pdf"
        assert body["ocr"] is True
        doc_id = body["document"]["doc_id"]
        CREATED_DOCS.append(doc_id)
        final = _poll_doc(auth_headers, doc_id, timeout=180)
        assert final["status"] == "ready", f"err={final.get('error')}"
        assert final.get("pages") == 2, f"pages={final.get('pages')}"


# ---------- 6. Blank PNG → error ----------

class TestKbBlankPng:
    def test_blank_png_errors(self, auth_headers):
        raw = _gen_blank_png()
        payload = {"title": "TEST_blank", "filename": "blank.png", "mime": "image/png", "data_base64": _b64(raw)}
        r = requests.post(f"{API}/kb/documents", headers=auth_headers, json=payload, timeout=20)
        # Should be 200 ingested → ocr pipeline runs → eventually error
        assert r.status_code == 200, r.text
        doc_id = r.json()["document"]["doc_id"]
        CREATED_DOCS.append(doc_id)
        final = _poll_doc(auth_headers, doc_id, timeout=120)
        assert final["status"] == "error", f"expected error, got {final}"
        assert "OCR" in (final.get("error") or "") or "leggibile" in (final.get("error") or "").lower()


# ---------- 7. Text-layer PDF ----------

class TestKbTextPdf:
    def test_text_pdf(self, auth_headers):
        txt = "Documento di prova con testo nativo non scansionato. " * 20
        raw = _gen_text_pdf(txt)
        payload = {"title": "TEST_text_pdf", "filename": "text.pdf", "mime": "application/pdf", "data_base64": _b64(raw)}
        r = requests.post(f"{API}/kb/documents", headers=auth_headers, json=payload, timeout=20)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["document"]["source_type"] == "pdf"
        assert body["ocr"] is False, f"should not need OCR, body={body}"
        doc_id = body["document"]["doc_id"]
        CREATED_DOCS.append(doc_id)
        final = _poll_doc(auth_headers, doc_id, timeout=60)
        assert final["status"] == "ready", f"err={final.get('error')}"


# ---------- 8. Edge cases: >20MB and random garbage ----------

class TestKbEdgeCases:
    def test_oversize_413(self, auth_headers):
        # ~21MB
        raw = os.urandom(21 * 1024 * 1024)
        payload = {"title": "TEST_oversize", "filename": "big.bin", "mime": "application/octet-stream", "data_base64": _b64(raw)}
        r = requests.post(f"{API}/kb/documents", headers=auth_headers, json=payload, timeout=30)
        assert r.status_code == 413, f"expected 413, got {r.status_code} {r.text[:200]}"

    def test_binary_garbage(self, auth_headers):
        raw = os.urandom(2048)
        payload = {"title": "TEST_garbage", "filename": "garbage.bin", "mime": "application/octet-stream", "data_base64": _b64(raw)}
        r = requests.post(f"{API}/kb/documents", headers=auth_headers, json=payload, timeout=20)
        # must not be 500: either 400, or 200 (treated as text fallback)
        assert r.status_code != 500, f"got 500: {r.text[:300]}"
        assert r.status_code in (200, 400), f"got unexpected {r.status_code}"
        if r.status_code == 200:
            doc_id = r.json()["document"]["doc_id"]
            CREATED_DOCS.append(doc_id)


# ---------- 9. OG preview ----------

class TestOgPreview:
    def test_og_png_public(self):
        r = requests.get(f"{API}/public/og-preview.png", timeout=30)
        assert r.status_code == 200, r.text
        assert r.headers.get("content-type", "").startswith("image/png")
        assert "cache-control" in {k.lower() for k in r.headers.keys()}
        assert len(r.content) > 1000
        # verify dimensions
        from PIL import Image
        img = Image.open(io.BytesIO(r.content))
        assert img.size == (1200, 630), f"size={img.size}"

    def test_cms_og_no_auth(self):
        r = requests.get(f"{API}/landing/cms/og", timeout=10)
        assert r.status_code in (401, 403), f"got {r.status_code}"

    def test_cms_og_with_auth(self, auth_headers):
        r = requests.get(f"{API}/landing/cms/og", headers=auth_headers, timeout=10)
        assert r.status_code == 200, r.text
        data = r.json()
        for k in ("generated_at", "left", "limit", "exists"):
            assert k in data, f"missing key {k} in {data}"

    def test_cms_og_refresh_wrong_password(self, auth_headers):
        h = {**auth_headers, "X-CMS-Password": "wrong_pw"}
        r = requests.post(f"{API}/landing/cms/og/refresh", headers=h, timeout=30)
        assert r.status_code in (401, 403), f"got {r.status_code}"

    def test_cms_og_refresh_correct(self, auth_headers):
        h = {**auth_headers, "X-CMS-Password": CMS_PASSWORD}
        r = requests.post(f"{API}/landing/cms/og/refresh", headers=h, timeout=60)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d.get("ok") is True
        assert isinstance(d.get("left"), int)
        assert isinstance(d.get("limit"), int)


# ---------- 10. FAQ semantic suggest ----------

class TestFaqSuggest:
    def test_faq_regression(self):
        r = requests.get(f"{API}/public/faq", timeout=10)
        assert r.status_code == 200
        assert len(r.json().get("items", [])) >= 1

    def test_suggest_crediti(self):
        r = requests.get(f"{API}/public/faq/suggest", params={"q": "ho finito i crediti", "k": 3}, timeout=20)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d.get("engine") == "minilm-local", f"engine={d.get('engine')}"
        items = d.get("items", [])
        assert 1 <= len(items) <= 3, f"items count={len(items)}"
        first_q = (items[0].get("q") or "").lower()
        assert "credit" in first_q, f"first q doesn't mention crediti: {first_q}"

    def test_suggest_short_query(self):
        r = requests.get(f"{API}/public/faq/suggest", params={"q": "ab"}, timeout=10)
        assert r.status_code == 200
        d = r.json()
        assert d.get("items") == []
        assert d.get("engine") == "none"

    def test_suggest_k_clamped(self):
        r = requests.get(f"{API}/public/faq/suggest", params={"q": "come carico contratto esempio axon", "k": 50}, timeout=20)
        assert r.status_code == 200
        d = r.json()
        assert len(d.get("items", [])) <= 6, f"k not clamped: {len(d.get('items'))}"


# ---------- Cleanup ----------

class TestCleanup:
    def test_zz_cleanup(self, auth_headers):
        """Delete all docs created in this run."""
        for doc_id in list(CREATED_DOCS):
            r = requests.delete(f"{API}/kb/documents/{doc_id}", headers=auth_headers, timeout=10)
            # 200 or 404 both ok
            assert r.status_code in (200, 404), f"delete {doc_id} → {r.status_code}"
