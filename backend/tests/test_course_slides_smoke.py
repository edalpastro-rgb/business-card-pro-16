"""Smoke test Fase 3 Area Corsi — slide PowerPoint (manuale + AI + pptx + attach). Esegui: python tests/test_course_slides_smoke.py"""
import base64
import io
import os
import sys

import requests

BASE = os.environ.get("BASE_URL", "http://localhost:8001/api")


def login(email, pwd):
    r = requests.post(f"{BASE}/auth/login", json={"email": email, "password": pwd}, timeout=30)
    r.raise_for_status()
    return {"Authorization": f"Bearer {r.json()['token']}"}


def main():
    h = login("test@dalpastro.it", "Test1234!")
    hs = login("sub1@dalpastro.it", "Sub12345!")
    body_text = ("Il fotovoltaico residenziale: dimensionamento impianto in base ai consumi annui (kWh), scelta inverter, "
                 "accumulo con batterie al litio, detrazioni fiscali 50%, tempi di rientro 6-8 anni, gestione obiezioni del cliente "
                 "su costi iniziali e manutenzione, comunità energetiche rinnovabili e incentivi GSE sull'energia condivisa.")
    c = requests.post(f"{BASE}/courses", json={"title": "Corso Fotovoltaico Smoke", "description": "Formazione consulenti", "body_text": body_text,
                                              "files": [{"name": "note.txt", "mime": "text/plain", "data_base64": base64.b64encode(b"Slide extra: normativa CER 2024, autoconsumo diffuso.").decode()}]}, headers=h, timeout=30)
    assert c.status_code == 200, c.text
    cid = c.json()["course_id"]
    try:
        # vuoto
        r = requests.get(f"{BASE}/courses/{cid}/slides", headers=h, timeout=30); assert r.json()["slides"] is None and r.json()["can_edit"], r.text
        # manuale
        r = requests.put(f"{BASE}/courses/{cid}/slides", json={"title": "Deck manuale", "subtitle": "test", "slides": [{"title": "Agenda", "bullets": ["Punto 1", "Punto 2"], "notes": "nota"}, {"title": "", "bullets": [], "notes": ""}]}, headers=h, timeout=30)
        assert r.status_code == 200 and len(r.json()["slides"]["slides"]) == 1 and r.json()["slides"]["source"] == "manual", r.text
        # collaboratore: non visibile
        r = requests.get(f"{BASE}/courses/{cid}/slides", headers=hs, timeout=30); assert r.json()["slides"] is None and not r.json()["can_edit"], r.text
        r = requests.get(f"{BASE}/courses/{cid}/slides/pptx", headers=hs, timeout=30); assert r.status_code == 404, r.text
        r = requests.put(f"{BASE}/courses/{cid}/slides", json={"title": "x"}, headers=hs, timeout=30); assert r.status_code == 403, r.text
        # pptx creatore
        r = requests.get(f"{BASE}/courses/{cid}/slides/pptx", headers=h, timeout=60); assert r.status_code == 200, r.text
        raw = base64.b64decode(r.json()["data_base64"])
        from pptx import Presentation
        prs = Presentation(io.BytesIO(raw)); assert len(prs.slides) == 3, len(prs.slides)  # copertina + 1 + chiusura
        print("pptx ok:", r.json()["name"], len(raw), "bytes,", len(prs.slides), "slides")
        # visibile → collaboratore scarica
        r = requests.put(f"{BASE}/courses/{cid}/slides", json={"visible_to_recipients": True}, headers=h, timeout=30); assert r.json()["slides"]["visible_to_recipients"] is True
        r = requests.get(f"{BASE}/courses/{cid}/slides/pptx", headers=hs, timeout=60); assert r.status_code == 200, r.text
        # attach
        r = requests.post(f"{BASE}/courses/{cid}/slides/attach", headers=h, timeout=60); assert r.status_code == 200, r.text
        r = requests.post(f"{BASE}/courses/{cid}/slides/attach", headers=h, timeout=60); assert r.status_code == 200, r.text
        files = requests.get(f"{BASE}/courses/{cid}", headers=h, timeout=30).json()["files"]
        assert sum(1 for f in files if f["name"].endswith(".pptx")) == 1, files  # sostituisce, non duplica
        print("attach ok:", [f["name"] for f in files])
        # AI
        r = requests.post(f"{BASE}/courses/{cid}/slides/generate", json={"n_slides": 5, "instructions": "tono semplice"}, headers=h, timeout=180)
        assert r.status_code == 200, r.text
        d = r.json()["slides"]
        assert d["source"] == "ai" and d["visible_to_recipients"] is True and 3 <= len(d["slides"]) <= 20, d
        print("AI ok:", d["title"], "|", [s["title"] for s in d["slides"]], "| materiali:", r.json()["materials_chars"], "chars")
        # AI then manual edit → ai+manual
        r = requests.put(f"{BASE}/courses/{cid}/slides", json={"slides": d["slides"][:2]}, headers=h, timeout=30); assert r.json()["slides"]["source"] == "ai+manual", r.text
        r = requests.delete(f"{BASE}/courses/{cid}/slides", headers=h, timeout=30); assert r.status_code == 200
        assert requests.get(f"{BASE}/courses/{cid}/slides", headers=h, timeout=30).json()["slides"] is None
        print("ALL OK")
    finally:
        requests.delete(f"{BASE}/courses/{cid}", headers=h, timeout=30)


if __name__ == "__main__":
    sys.exit(main())
